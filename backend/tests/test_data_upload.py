"""The Data Upload Center: templates, validation, import, history and rollback.

Test files are small and deterministic, built in-memory from the codes the Phase
2 fixture seeds. No production file is read and no fake production data is
created: the rows here reference the test warehouse's own master records.
"""

from __future__ import annotations

import csv
import io
import time

import pytest
from fastapi.testclient import TestClient
from sqlalchemy import func, select
from sqlalchemy.orm import Session

from app.api.deps import get_session
from app.auth.security import hash_password
from app.database.models import DimRegion
from app.database.models_admin import (
    ImportMode,
    UploadBatch,
    UploadError,
    UploadStatus,
)
from app.database.models_ai import AppUser, AuditAction, AuditLog, Role, UserStatus
from app.database.models_warehouse import FactSales
from app.main import app
from app.security.sections import ALLOW, DENY, SectionKey
from app.upload.registry import UPLOAD_TYPE_BY_KEY, get_upload_type
from app.upload.templates import build_csv, build_xlsx

pytest.importorskip("multipart", reason="python-multipart is required by FastAPI forms")

PASSWORD = "Correct-Horse-9"

XLSX_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


@pytest.fixture
def upload_client(agent_engine, users, monkeypatch, tmp_path):
    """A client signed in as a super administrator, with staging in tmp_path."""
    import app.database.connection as connection
    import app.upload.files as upload_files

    # The ETL opens its own session from the global engine; point it at the test
    # warehouse so an upload lands in the fixture's database, not a real one.
    #
    # ``service`` is patched by name as well because it bound ``get_engine`` at
    # import time. ``jobs`` deliberately does not — it calls
    # ``connection.get_engine()`` through the module precisely so that patching
    # the module reaches the worker threads too.
    monkeypatch.setattr(connection, "get_engine", lambda *a, **k: agent_engine)
    monkeypatch.setattr("app.upload.service.get_engine", lambda *a, **k: agent_engine)
    monkeypatch.setattr(upload_files, "upload_dir", lambda: tmp_path)

    with Session(agent_engine) as session:
        session.add(AppUser(username="root", display_name="Administrator",
                            role=Role.SUPER_ADMIN, is_active=True,
                            status=UserStatus.ACTIVE,
                            password_hash=hash_password(PASSWORD)))
        for user in session.query(AppUser).all():
            if user.password_hash is None:
                user.password_hash = hash_password(PASSWORD)
        session.commit()

    def _session_override():
        session = Session(bind=agent_engine, expire_on_commit=False, future=True)
        try:
            yield session
        finally:
            session.close()

    app.dependency_overrides[get_session] = _session_override
    try:
        yield TestClient(app)
    finally:
        app.dependency_overrides.clear()


def login(client: TestClient, username: str = "root") -> str:
    response = client.post("/api/auth/login",
                           json={"username": username, "password": PASSWORD})
    assert response.status_code == 200, response.text
    return response.json()["access_token"]


def auth(token: str) -> dict[str, str]:
    return {"Authorization": f"Bearer {token}"}


def csv_bytes(headers: list[str], rows: list[list[object]]) -> bytes:
    buffer = io.StringIO(newline="")
    writer = csv.writer(buffer, lineterminator="\r\n")
    writer.writerow(headers)
    writer.writerows(rows)
    return buffer.getvalue().encode("utf-8")


class FinishedUpload:
    """A finished job, in the shape ``POST /preview`` used to answer with.

    Uploading is now a background job: the request returns ``202`` and a job to
    watch, and the counts, preview rows and errors land on the batch afterwards.
    Almost every test here is about *what the validation decided*, not about how
    the result was delivered, so the wait-then-read is done once here and the
    assertions stay as they were.
    """

    def __init__(self, status_code: int, payload: dict):
        self.status_code = status_code
        self._payload = payload

    def json(self) -> dict:
        return self._payload

    @property
    def text(self) -> str:
        return str(self._payload)


def wait_for_job(client: TestClient, token: str, upload_id: int,
                 timeout: float = 60.0) -> dict:
    """Poll until the batch reaches a terminal status, then return its detail.

    Polling rather than reaching into the executor, because that is exactly what
    the browser does — a test that waited on the future would be exercising a
    path no client can use.
    """
    deadline = time.monotonic() + timeout
    while True:
        detail = client.get(f"/api/data-upload/history/{upload_id}",
                            headers=auth(token)).json()
        status_value = detail["upload"]["status"]
        if status_value in UploadStatus.TERMINAL or status_value == UploadStatus.VALIDATED:
            return detail
        if time.monotonic() > deadline:
            raise AssertionError(
                f"upload {upload_id} was still {status_value} after {timeout}s"
            )
        time.sleep(0.02)


def _as_outcome(detail: dict) -> dict:
    """The old ``/preview`` response, rebuilt from the batch it produced.

    The error entries are already identical — ``UploadIssue.to_dict`` and the
    history endpoint emit the same keys — so only the envelope differs.
    """
    return {
        "upload": detail["upload"],
        "preview": detail.get("preview"),
        "warnings": detail.get("warnings", []),
        "errors": detail.get("errors", []),
        "error_count": detail.get("error_total", 0),
        "error_truncated": False,
    }


def upload(client: TestClient, token: str, upload_type: str, content: bytes,
           filename: str = "test.csv", mode: str = ImportMode.UPSERT,
           content_type: str = "text/csv"):
    """Upload a file and wait for its validation job to finish.

    A rejection (bad extension, oversized, unknown type) is returned as the real
    response, because there is no job to wait for and the test is asserting on
    the refusal itself.
    """
    response = client.post(
        "/api/data-upload/preview", headers=auth(token),
        files={"file": (filename, content, content_type)},
        data={"upload_type": upload_type, "import_mode": mode},
    )
    if response.status_code != 202:
        return response
    upload_id = response.json()["upload"]["upload_id"]
    return FinishedUpload(202, _as_outcome(wait_for_job(client, token, upload_id)))


def commit(client: TestClient, token: str, upload_id: int, confirm: bool = True):
    """Confirm an import and wait for it. Mirrors :func:`upload`."""
    response = client.post(f"/api/data-upload/{upload_id}/commit",
                           params={"confirm": confirm}, headers=auth(token))
    if response.status_code != 202:
        return response
    return FinishedUpload(202, _as_outcome(wait_for_job(client, token, upload_id)))


# ==========================================================================
# Catalogue and templates
# ==========================================================================


def test_the_catalogue_is_derived_from_the_real_schema(upload_client):
    token = login(upload_client)
    body = upload_client.get("/api/data-upload/types", headers=auth(token)).json()

    master = next(c for c in body["categories"] if c["key"] == "MASTER")
    transactional = next(c for c in body["categories"] if c["key"] == "TRANSACTIONAL")
    master_keys = {t["key"] for t in master["types"]}
    transactional_keys = {t["key"] for t in transactional["types"]}

    # The Phase 1 hierarchy and product dimension, not invented tables.
    assert {"dim_company", "dim_business_unit", "dim_sales_line", "dim_zone",
            "dim_region", "dim_area", "dim_unit", "dim_territory",
            "dim_sub_territory", "dim_material"} <= master_keys
    # The Phase 2 fact data types, and only those. Collection and Outstanding
    # went with their datasets in revision 0020: the catalogue is derived from
    # `etl.datasets.DATASETS`, so an upload type cannot outlive its pipeline.
    assert transactional_keys == {"sales", "material_stock", "target"}
    assert body["limits"]["extensions"] == [".xlsx", ".csv"]


def test_every_upload_type_declares_a_business_key_and_required_columns():
    for upload_type in UPLOAD_TYPE_BY_KEY.values():
        assert upload_type.business_key, upload_type.key
        assert upload_type.business_key_description, upload_type.key
        assert upload_type.required_columns, upload_type.key


def test_the_csv_template_has_the_headers_the_validator_expects(upload_client):
    token = login(upload_client)
    response = upload_client.get("/api/data-upload/types/dim_region/template",
                                 params={"format": "csv"}, headers=auth(token))
    assert response.status_code == 200
    assert "attachment" in response.headers["content-disposition"]

    rows = list(csv.reader(response.content.decode("utf-8-sig").splitlines()))
    spec = get_upload_type("dim_region")
    assert rows[0] == [column.name for column in spec.columns]
    assert len(rows) == 2       # header plus exactly one example row


def test_the_xlsx_template_carries_an_instruction_sheet():
    from openpyxl import load_workbook

    workbook = load_workbook(io.BytesIO(build_xlsx(get_upload_type("sales"))))
    assert workbook.sheetnames == ["Data", "Instructions"]
    text = "\n".join(
        str(cell.value) for row in workbook["Instructions"].iter_rows()
        for cell in row if cell.value
    )
    assert "REQUIRED" in text
    assert "Record identity" in text or "invoice_no" in text


def test_a_template_contains_no_real_business_data(upload_client, agent_engine):
    """A template is a public artefact; it must not leak the warehouse.

    Proved by putting a value in the warehouse that no schema example could
    coincidentally match, then checking it never reaches a template. The example
    row is a static placeholder declared in ``master_data.schema``; nothing in
    template generation reads a table.
    """
    marker = "Zzz-Confidential-Region-9174"
    with Session(agent_engine) as session:
        session.add(DimRegion(region_code="REG-SECRET", region_name=marker,
                              zone_code="Z001", region_head_name=marker))
        session.commit()

    for upload_type in ("dim_region", "dim_material", "sales"):
        content = build_csv(get_upload_type(upload_type)).decode("utf-8-sig")
        assert marker not in content, upload_type
        assert "REG-SECRET" not in content, upload_type


def test_a_template_download_is_audited(upload_client, agent_engine):
    token = login(upload_client)
    upload_client.get("/api/data-upload/types/sales/template", headers=auth(token))
    with Session(agent_engine) as session:
        entry = session.execute(
            select(AuditLog).where(AuditLog.action == AuditAction.TEMPLATE_DOWNLOADED)
        ).scalars().first()
    assert entry is not None
    assert entry.resource == "template:sales"


# ==========================================================================
# File validation
# ==========================================================================


def test_an_unsupported_extension_is_refused(upload_client):
    token = login(upload_client)
    response = upload_client.post(
        "/api/data-upload/preview", headers=auth(token),
        files={"file": ("payload.exe", b"MZ\x90\x00", "application/octet-stream")},
        data={"upload_type": "dim_region", "import_mode": ImportMode.UPSERT},
    )
    assert response.status_code == 415
    assert ".xlsx" in response.json()["detail"]


def test_a_renamed_binary_is_refused_even_with_a_csv_extension(upload_client):
    """The extension and the MIME type both say CSV; the bytes do not."""
    token = login(upload_client)
    response = upload(upload_client, token, "dim_region", b"MZ\x90\x00binary",
                      filename="evil.csv")
    assert response.status_code == 415
    assert "binary" in response.json()["detail"].lower()


def test_an_empty_file_is_refused(upload_client):
    token = login(upload_client)
    response = upload(upload_client, token, "dim_region", b"", filename="empty.csv")
    assert response.status_code == 415


def test_an_oversized_file_is_refused(upload_client, monkeypatch):
    import app.upload.files as upload_files

    monkeypatch.setattr(upload_files, "MAX_UPLOAD_BYTES", 64)
    token = login(upload_client)
    response = upload(upload_client, token, "dim_region", b"a" * 500,
                      filename="big.csv")
    assert response.status_code == 415
    assert "larger than" in response.json()["detail"]


def test_an_unknown_upload_type_is_a_404(upload_client):
    token = login(upload_client)
    response = upload(upload_client, token, "dim_unicorn", csv_bytes(["A"], [["1"]]))
    assert response.status_code == 404


def test_a_file_missing_a_required_column_loads_nothing(upload_client):
    token = login(upload_client)
    response = upload(upload_client, token, "dim_region",
                      csv_bytes(["Region Code"], [["REG900"]]))
    # 202: the file was accepted and queued. Whether it validates is the job's
    # verdict, not the request's, and this file's is that a column is missing.
    assert response.status_code == 202
    body = response.json()
    assert body["upload"]["status"] == UploadStatus.FAILED
    assert body["upload"]["totals"]["valid_rows"] == 0
    assert any(issue["error_code"] == "MISSING_COLUMN" for issue in body["errors"])


# ==========================================================================
# Master data upload
# ==========================================================================


MASTER_HEADERS = ["Zone Code", "Region Code", "Region", "Region Head ID",
                  "Region Head", "Region HQ", "Location"]


def test_valid_master_data_is_previewed_then_imported(upload_client, agent_engine):
    token = login(upload_client)
    content = csv_bytes(MASTER_HEADERS, [
        ["Z001", "REG900", "Sylhet", "EMP900", "Md. Nabil", "Sylhet", "Sylhet North"],
        ["Z001", "REG901", "Barishal", "EMP901", "Md. Sabbir", "Barishal", "Barishal"],
    ])

    preview = upload(upload_client, token, "dim_region", content).json()
    assert preview["upload"]["status"] == UploadStatus.VALIDATED
    assert preview["upload"]["totals"]["total_rows"] == 2
    assert preview["upload"]["totals"]["valid_rows"] == 2
    assert preview["upload"]["totals"]["invalid_rows"] == 0
    assert preview["upload"]["can_commit"] is True
    assert len(preview["preview"]["rows"]) == 2

    # Nothing is written before the user confirms.
    with Session(agent_engine) as session:
        assert session.execute(
            select(func.count()).select_from(DimRegion)
            .where(DimRegion.region_code == "REG900")
        ).scalar_one() == 0

    upload_id = preview["upload"]["upload_id"]
    result = commit(upload_client, token, upload_id)
    assert result.status_code == 202, result.text
    body = result.json()
    assert body["upload"]["status"] == UploadStatus.COMPLETED
    assert body["upload"]["totals"]["inserted_rows"] == 2

    with Session(agent_engine) as session:
        region = session.execute(
            select(DimRegion).where(DimRegion.region_code == "REG900")
        ).scalar_one()
    assert region.region_name == "Sylhet"
    assert region.zone_code == "Z001"


def test_importing_without_confirmation_is_refused(upload_client):
    token = login(upload_client)
    preview = upload(upload_client, token, "dim_region", csv_bytes(
        MASTER_HEADERS,
        [["Z001", "REG902", "Rangpur", "", "", "", ""]],
    )).json()
    response = upload_client.post(
        f"/api/data-upload/{preview['upload']['upload_id']}/commit",
        headers=auth(token),
    )
    assert response.status_code == 400
    assert "Confirm" in response.json()["detail"]


def test_invalid_master_data_reports_the_row_column_and_a_fix(upload_client):
    token = login(upload_client)
    content = csv_bytes(MASTER_HEADERS, [
        ["Z001", "REG903", "Cumilla", "", "", "", ""],           # valid
        ["Z999", "REG904", "Nowhere", "", "", "", ""],           # unknown parent
        ["Z001", "", "Nameless", "", "", "", ""],                # missing key
    ])
    body = upload(upload_client, token, "dim_region", content).json()

    assert body["upload"]["totals"]["total_rows"] == 3
    assert body["upload"]["totals"]["valid_rows"] == 1
    assert body["upload"]["totals"]["invalid_rows"] == 2

    by_code = {issue["error_code"] for issue in body["errors"]}
    assert "INVALID_PARENT_CODE" in by_code
    assert "MISSING_REQUIRED_FIELD" in by_code
    for issue in body["errors"]:
        assert issue["row"] is not None
        assert issue["column"]
        assert issue["suggested_fix"]


def test_duplicate_master_codes_within_one_file_are_detected(upload_client):
    token = login(upload_client)
    content = csv_bytes(MASTER_HEADERS, [
        ["Z001", "REG905", "Tangail", "", "", "", ""],
        ["Z001", "REG905", "Tangail again", "", "", "", ""],
    ])
    body = upload(upload_client, token, "dim_region", content).json()
    assert body["upload"]["totals"]["duplicate_rows"] == 1
    assert any(i["error_code"] == "DUPLICATE_IN_FILE" for i in body["errors"])


def test_insert_mode_refuses_a_code_that_already_exists(upload_client):
    token = login(upload_client)
    content = csv_bytes(MASTER_HEADERS,
                        [["Z001", "REG001", "Dhaka renamed", "", "", "", ""]])
    body = upload(upload_client, token, "dim_region", content,
                  mode=ImportMode.INSERT).json()
    assert body["upload"]["status"] == UploadStatus.FAILED
    assert any(i["error_code"] == "RECORD_ALREADY_EXISTS" for i in body["errors"])


def test_update_mode_refuses_a_code_that_does_not_exist_yet(upload_client):
    token = login(upload_client)
    content = csv_bytes(MASTER_HEADERS,
                        [["Z001", "REG906", "Brand new", "", "", "", ""]])
    body = upload(upload_client, token, "dim_region", content,
                  mode=ImportMode.UPDATE).json()
    assert any(i["error_code"] == "RECORD_NOT_FOUND" for i in body["errors"])


def test_upsert_updates_an_existing_master_record_in_place(upload_client, agent_engine):
    token = login(upload_client)
    with Session(agent_engine) as session:
        before = session.execute(
            select(func.count()).select_from(DimRegion)
        ).scalar_one()

    content = csv_bytes(MASTER_HEADERS,
                        [["Z001", "REG001", "Dhaka Metro", "", "", "Dhaka", ""]])
    preview = upload(upload_client, token, "dim_region", content).json()
    commit(upload_client, token, preview['upload']['upload_id'])

    with Session(agent_engine) as session:
        after = session.execute(select(func.count()).select_from(DimRegion)).scalar_one()
        region = session.execute(
            select(DimRegion).where(DimRegion.region_code == "REG001")
        ).scalar_one()
    # Updated, never duplicated.
    assert after == before
    assert region.region_name == "Dhaka Metro"


def test_a_master_upload_never_deletes_a_record_the_file_omits(upload_client,
                                                               agent_engine):
    token = login(upload_client)
    with Session(agent_engine) as session:
        codes_before = {
            code for (code,) in session.execute(select(DimRegion.region_code)).all()
        }

    content = csv_bytes(MASTER_HEADERS,
                        [["Z001", "REG907", "Only row", "", "", "", ""]])
    preview = upload(upload_client, token, "dim_region", content).json()
    commit(upload_client, token, preview['upload']['upload_id'])

    with Session(agent_engine) as session:
        codes_after = {
            code for (code,) in session.execute(select(DimRegion.region_code)).all()
        }
    assert codes_before < codes_after       # strictly grown; nothing removed


# ==========================================================================
# Transactional upload
# ==========================================================================


#: Net Sales is the required measure, so an upload fixture has to carry it.
#: Gross is kept as well, because a real file usually sends the breakdown.
SALES_HEADERS = ["Date", "Invoice No", "SKU Code", "Territory Code", "Quantity",
                 "Gross Sales", "Discount", "Net Sales"]


def sales_rows(*invoices: str) -> list[list[object]]:
    return [
        ["2026-08-20", invoice, "SKU001", "TR001", 10, 12000, 1000, 11000]
        for invoice in invoices
    ]


def test_valid_transaction_data_reaches_the_fact_table(upload_client, agent_engine):
    token = login(upload_client)
    content = csv_bytes(SALES_HEADERS, sales_rows("UPL-001", "UPL-002"))

    preview = upload(upload_client, token, "sales", content).json()
    assert preview["upload"]["status"] == UploadStatus.VALIDATED
    assert preview["upload"]["totals"]["valid_rows"] == 2

    # A dry run must leave the warehouse untouched.
    with Session(agent_engine) as session:
        assert session.execute(
            select(func.count()).select_from(FactSales)
            .where(FactSales.invoice_no == "UPL-001")
        ).scalar_one() == 0

    result = commit(upload_client, token, preview['upload']['upload_id']).json()
    assert result["upload"]["status"] == UploadStatus.COMPLETED
    assert result["upload"]["totals"]["inserted_rows"] == 2
    assert result["upload"]["etl_batch_id"] is not None

    with Session(agent_engine) as session:
        row = session.execute(
            select(FactSales).where(FactSales.invoice_no == "UPL-001")
        ).scalar_one()
    # The Phase 2 pipeline resolved the hierarchy from the territory code alone.
    assert row.region_id is not None
    assert row.source_system == "UPLOAD"
    assert float(row.net_sales) == 11000


def test_a_duplicate_transaction_is_detected_by_its_business_key(upload_client,
                                                                 agent_engine):
    token = login(upload_client)
    content = csv_bytes(SALES_HEADERS, sales_rows("UPL-010"))

    first = upload(upload_client, token, "sales", content).json()
    commit(upload_client, token, first['upload']['upload_id'])

    # The same file again: the business key (invoice + SKU + source) already exists.
    second = upload(upload_client, token, "sales", content).json()
    assert second["upload"]["totals"]["duplicate_rows"] == 1
    assert any("already in the warehouse" in warning
               for warning in second["warnings"])

    commit(upload_client, token, second['upload']['upload_id'])

    with Session(agent_engine) as session:
        count = session.execute(
            select(func.count()).select_from(FactSales)
            .where(FactSales.invoice_no == "UPL-010")
        ).scalar_one()
    # Updated in place — the transaction exists exactly once.
    assert count == 1


def test_duplicates_inside_one_file_are_reported(upload_client):
    token = login(upload_client)
    content = csv_bytes(SALES_HEADERS, sales_rows("UPL-020", "UPL-020"))
    body = upload(upload_client, token, "sales", content).json()
    assert any(i["error_code"] == "DUPLICATE_IN_FILE" for i in body["errors"])


def test_invalid_master_mapping_is_rejected_with_a_suggested_fix(upload_client):
    token = login(upload_client)
    content = csv_bytes(SALES_HEADERS, [
        ["2026-08-20", "UPL-030", "SKU-NOPE", "TR001", 1, 100, 0, 100],
        ["2026-08-20", "UPL-031", "SKU001", "TR-NOPE", 1, 100, 0, 100],
        ["2026-08-20", "UPL-032", "SKU001", "TR001", 1, 100, 0, 100],
    ])
    body = upload(upload_client, token, "sales", content).json()

    assert body["upload"]["totals"]["valid_rows"] == 1
    assert body["upload"]["totals"]["invalid_rows"] == 2
    codes = {i["error_code"] for i in body["errors"]}
    assert "INVALID_MATERIAL_CODE" in codes
    assert "INVALID_TERRITORY_CODE" in codes
    for issue in body["errors"]:
        assert issue["suggested_fix"], issue


def test_an_invalid_date_is_rejected_not_guessed(upload_client):
    token = login(upload_client)
    content = csv_bytes(SALES_HEADERS,
                        [["not-a-date", "UPL-040", "SKU001", "TR001", 1, 100, 0]])
    body = upload(upload_client, token, "sales", content).json()
    assert body["upload"]["totals"]["invalid_rows"] == 1
    assert any(i["error_code"].startswith("INVALID_DATE")
               or i["error_code"] == "AMBIGUOUS_DATE" for i in body["errors"])


def test_a_row_without_an_organisational_code_is_rejected(upload_client):
    token = login(upload_client)
    content = csv_bytes(["Date", "Invoice No", "SKU Code", "Quantity", "Net Sales"],
                        [["2026-08-20", "UPL-050", "SKU001", 1, 100]])
    body = upload(upload_client, token, "sales", content).json()
    assert any(i["error_code"] == "NO_ORGANISATIONAL_CODE" for i in body["errors"])


# ==========================================================================
# Batch-aware lines: preview, duplicate report and validation counters
# ==========================================================================


BATCH_HEADERS = [*SALES_HEADERS, "Batch Code"]


def batch_rows(*batches: str, invoice: str = "UPL-100") -> list[list[object]]:
    return [
        ["2026-08-20", invoice, "SKU001", "TR001", 10, 12000, 1000, 11000, batch]
        for batch in batches
    ]


def test_the_preview_names_the_line_and_its_verdict(upload_client):
    """Item 10: row, identity, volume and reason, before anything is committed."""
    token = login(upload_client)
    content = csv_bytes(BATCH_HEADERS, batch_rows("BATCH-A", "BATCH-B", "BATCH-A"))
    body = upload(upload_client, token, "sales", content).json()

    columns = body["preview"]["columns"]
    for name in ("__row", "__invoice_no", "__material_code", "__batch_code",
                 "__quantity", "__volume", "__reason"):
        assert name in columns, name
    # No unit column beside the volume: the file states one total per line.
    assert "__volume_unit" not in columns

    rows = body["preview"]["rows"]
    assert [row["__batch_code"] for row in rows] == ["BATCH-A", "BATCH-B", "BATCH-A"]
    assert [row["__status"] for row in rows] == ["VALID", "VALID", "INVALID"]
    assert rows[0]["__reason"] == "Valid"
    assert "Duplicate line found within uploaded file" in rows[2]["__reason"]


def test_two_batches_on_one_invoice_both_load(upload_client, agent_engine):
    """The headline rule, end to end through the upload centre."""
    token = login(upload_client)
    content = csv_bytes(BATCH_HEADERS,
                        batch_rows("BATCH-A", "BATCH-B", invoice="UPL-110"))

    preview = upload(upload_client, token, "sales", content).json()
    assert preview["upload"]["totals"]["valid_rows"] == 2

    commit(upload_client, token, preview['upload']['upload_id'])

    with Session(agent_engine) as session:
        batches = set(session.execute(
            select(FactSales.batch_code).where(FactSales.invoice_no == "UPL-110")
        ).scalars())
    assert batches == {"BATCH-A", "BATCH-B"}


def test_the_line_and_validation_reports_are_served(upload_client):
    """Items 36 and 37 over HTTP, including the route ordering they need."""
    token = login(upload_client)
    content = csv_bytes(BATCH_HEADERS,
                        batch_rows("BATCH-A", "BATCH-B", "BATCH-A",
                                   invoice="UPL-120"))
    preview = upload(upload_client, token, "sales", content).json()
    committed = commit(upload_client, token, preview['upload']['upload_id']).json()
    batch_id = committed["upload"]["etl_batch_id"]

    lines = upload_client.get(f"/api/data-quality/{batch_id}/lines",
                              headers=auth(token)).json()
    assert [r["status"] for r in lines["records"]] == [
        "ACCEPTED", "ACCEPTED", "REJECTED"]
    assert lines["records"][2]["error_code"] == "DUPLICATE_IN_FILE"

    rejected = upload_client.get(f"/api/data-quality/{batch_id}/lines",
                                 params={"status": "REJECTED"},
                                 headers=auth(token)).json()
    assert rejected["total"] == 1

    validation = upload_client.get(f"/api/data-quality/{batch_id}/validation",
                                   headers=auth(token)).json()
    assert validation["counts"]["total_lines"] == 3
    assert validation["counts"]["valid_lines"] == 2
    assert validation["counts"]["duplicate_lines"] == 1

    # The literal routes must not be swallowed by /data-quality/{batch_id}.
    for path in ("/api/data-quality/master-mapping",
                 "/api/data-quality/master-sources/status"):
        assert upload_client.get(path, headers=auth(token)).status_code == 200


def test_the_unit_catalogue_is_no_longer_published(upload_client):
    """It served the volume-unit filter, and both are gone."""
    assert upload_client.get("/api/units").status_code == 404


def test_a_volume_unit_parameter_is_not_a_filter_any_more(upload_client):
    """An unknown query parameter is ignored, as it is everywhere else here.

    The point is that it cannot narrow the report: a sales line records one
    Total Volume and no unit, so there is no KG subset to ask for.
    """
    token = login(upload_client)
    response = upload_client.get("/api/reports/sales",
                                 params={"volume_unit": "GALLON"},
                                 headers=auth(token))
    assert response.status_code == 200
    assert "volume_unit" not in response.json()["filters"]


# ==========================================================================
# History, error report and rollback
# ==========================================================================


def test_upload_history_lists_the_batch_with_its_user_and_counts(upload_client):
    token = login(upload_client)
    content = csv_bytes(SALES_HEADERS, sales_rows("UPL-060"))
    created = upload(upload_client, token, "sales", content).json()

    history = upload_client.get("/api/data-upload/history", headers=auth(token)).json()
    assert history["total"] >= 1
    batch = next(b for b in history["batches"]
                 if b["upload_id"] == created["upload"]["upload_id"])
    assert batch["username"] == "root"
    assert batch["upload_type"] == "sales"
    assert batch["file_name"] == "test.csv"
    assert batch["totals"]["total_rows"] == 1


def test_a_batch_reports_the_job_fields_the_upload_dock_reads(upload_client):
    """``upload_uuid`` is the Import Job ID; no second identifier exists.

    The dock, the history and the job poll all render the same payload, so these
    keys are a contract rather than an implementation detail — and ``job_id``
    being an alias of ``upload_uuid`` is what stops the two ever disagreeing.
    """
    token = login(upload_client)
    content = csv_bytes(SALES_HEADERS, sales_rows("UPL-061"))
    batch = upload(upload_client, token, "sales", content).json()["upload"]

    assert batch["job_id"] == batch["upload_uuid"]
    # The validation job has finished by the time the helper returns, so the
    # stored progress is complete — never NULL, which is what the server default
    # guarantees for the batches that predate this column.
    assert batch["progress_percent"] == 100
    assert batch["stage"] in ("COMPLETED", "FAILED")
    assert batch["cancelled_at"] is None
    assert batch["cancelled_by"] is None
    # A finished validation is not active and cannot be cancelled; the run is
    # over by the time the response exists.
    assert batch["is_active"] is False
    assert batch["can_cancel"] is False
    assert batch["can_commit"] is True
    # Duration is derived from the timestamps rather than stored, and is a real
    # number once the batch has an end — the naive/aware mismatch SQLite creates
    # between a stored ``started_at`` and an in-session ``completed_at`` is
    # resolved rather than raising.
    assert batch["duration_seconds"] is None or batch["duration_seconds"] >= 0


def test_the_history_filter_offers_the_new_job_statuses(upload_client):
    """The status list is published from ``UploadStatus.ALL``, so the UI follows."""
    token = login(upload_client)
    statuses = upload_client.get("/api/data-upload/history",
                                 headers=auth(token)).json()["statuses"]
    assert "QUEUED" in statuses
    assert "CANCELLED" in statuses
    # PARTIAL is this system's COMPLETED_WITH_ERRORS and keeps its stored name:
    # renaming it would rewrite what every historical batch says about itself.
    assert "PARTIAL" in statuses


def test_a_batch_detail_lists_its_errors(upload_client):
    token = login(upload_client)
    content = csv_bytes(SALES_HEADERS,
                        [["2026-08-20", "UPL-070", "SKU-NOPE", "TR001", 1, 100, 0]])
    created = upload(upload_client, token, "sales", content).json()

    detail = upload_client.get(
        f"/api/data-upload/history/{created['upload']['upload_id']}",
        headers=auth(token),
    ).json()
    assert detail["error_total"] >= 1
    assert detail["upload_type"]["key"] == "sales"
    assert detail["errors"][0]["suggested_fix"]


def test_the_error_report_downloads_as_csv_with_the_expected_columns(upload_client):
    token = login(upload_client)
    content = csv_bytes(SALES_HEADERS,
                        [["2026-08-20", "UPL-080", "SKU-NOPE", "TR001", 1, 100, 0]])
    created = upload(upload_client, token, "sales", content).json()

    response = upload_client.get(
        f"/api/data-upload/history/{created['upload']['upload_id']}/errors",
        headers=auth(token),
    )
    assert response.status_code == 200
    assert response.headers["content-type"].startswith("text/csv")
    rows = list(csv.reader(response.content.decode("utf-8-sig").splitlines()))
    assert rows[0] == ["Row Number", "Column", "Value", "Error Code", "Error",
                       "Suggested Fix", "Severity"]
    assert len(rows) >= 2


def test_failed_records_are_listed_across_uploads(upload_client):
    token = login(upload_client)
    upload(upload_client, token, "sales", csv_bytes(
        SALES_HEADERS, [["2026-08-20", "UPL-090", "SKU-NOPE", "TR001", 1, 100, 0]]))

    body = upload_client.get("/api/data-upload/failed-records",
                             headers=auth(token)).json()
    assert body["total"] >= 1
    record = body["records"][0]
    assert record["upload_type"] == "sales"
    assert record["uploaded_by"] == "root"
    assert record["suggested_fix"]


def test_rollback_removes_only_this_batchs_fact_rows(upload_client, agent_engine):
    token = login(upload_client)
    content = csv_bytes(SALES_HEADERS, sales_rows("UPL-100", "UPL-101"))
    created = upload(upload_client, token, "sales", content).json()
    upload_id = created["upload"]["upload_id"]
    commit(upload_client, token, upload_id)

    with Session(agent_engine) as session:
        total_before = session.execute(
            select(func.count()).select_from(FactSales)
        ).scalar_one()

    response = upload_client.post(
        f"/api/data-upload/history/{upload_id}/rollback",
        params={"confirm": True}, headers=auth(token),
    )
    assert response.status_code == 200
    assert response.json()["fact_rows_removed"] == 2
    assert response.json()["upload"]["status"] == UploadStatus.ROLLED_BACK

    with Session(agent_engine) as session:
        assert session.execute(
            select(func.count()).select_from(FactSales)
            .where(FactSales.invoice_no.in_(["UPL-100", "UPL-101"]))
        ).scalar_one() == 0
        # The seeded facts from other batches are untouched.
        assert session.execute(
            select(func.count()).select_from(FactSales)
        ).scalar_one() == total_before - 2


def test_rollback_requires_confirmation(upload_client):
    token = login(upload_client)
    created = upload(upload_client, token, "sales",
                     csv_bytes(SALES_HEADERS, sales_rows("UPL-110"))).json()
    upload_id = created["upload"]["upload_id"]
    commit(upload_client, token, upload_id)

    response = upload_client.post(f"/api/data-upload/history/{upload_id}/rollback",
                                  headers=auth(token))
    assert response.status_code == 400


def test_rollback_is_refused_to_a_plain_administrator(upload_client, agent_engine):
    token = login(upload_client)
    created = upload(upload_client, token, "sales",
                     csv_bytes(SALES_HEADERS, sales_rows("UPL-120"))).json()
    upload_id = created["upload"]["upload_id"]
    commit(upload_client, token, upload_id)

    with Session(agent_engine) as session:
        session.add(AppUser(username="plain_admin", role=Role.ADMIN, is_active=True,
                            status=UserStatus.ACTIVE,
                            password_hash=hash_password(PASSWORD)))
        session.commit()

    admin_token = login(upload_client, "plain_admin")
    response = upload_client.post(f"/api/data-upload/history/{upload_id}/rollback",
                                  params={"confirm": True}, headers=auth(admin_token))
    assert response.status_code == 403


def test_a_master_upload_cannot_be_rolled_back(upload_client):
    token = login(upload_client)
    created = upload(upload_client, token, "dim_region", csv_bytes(
        MASTER_HEADERS, [["Z001", "REG950", "Norollback", "", "", "", ""]])).json()
    upload_id = created["upload"]["upload_id"]
    commit(upload_client, token, upload_id)

    response = upload_client.post(f"/api/data-upload/history/{upload_id}/rollback",
                                  params={"confirm": True}, headers=auth(token))
    assert response.status_code == 409
    assert "master-data upload" in response.json()["detail"]


# ==========================================================================
# Auditing and authorisation
# ==========================================================================


def test_uploading_and_importing_are_both_audited(upload_client, agent_engine):
    token = login(upload_client)
    created = upload(upload_client, token, "sales",
                     csv_bytes(SALES_HEADERS, sales_rows("UPL-130"))).json()
    commit(upload_client, token, created['upload']['upload_id'])

    with Session(agent_engine) as session:
        actions = {
            row.action for row in session.execute(
                select(AuditLog).where(AuditLog.username == "root")
            ).scalars()
        }
    assert AuditAction.DATA_UPLOADED in actions
    assert AuditAction.DATA_IMPORTED in actions


def test_the_upload_centre_is_refused_without_the_data_upload_section(upload_client):
    token = login(upload_client, "dhaka_rm")
    for path in ("/api/data-upload/types", "/api/data-upload/history",
                 "/api/data-upload/summary"):
        assert upload_client.get(path, headers=auth(token)).status_code == 403, path


def test_the_legacy_import_endpoint_also_requires_the_section(upload_client):
    token = login(upload_client, "dhaka_rm")
    response = upload_client.post("/api/import/sales/records", headers=auth(token),
                                  json={"records": [{"Invoice No": "X"}]})
    assert response.status_code == 403

    anonymous = upload_client.post("/api/import/sales/records",
                                   json={"records": [{"Invoice No": "X"}]})
    assert anonymous.status_code == 401


def test_a_granted_user_can_reach_the_upload_centre(upload_client, agent_engine):
    """The section is configurable per user, not fixed to the admin role."""
    admin_token = login(upload_client)
    with Session(agent_engine) as session:
        user_id = session.execute(
            select(AppUser.user_id).where(AppUser.username == "dhaka_rm")
        ).scalar_one()

    upload_client.put(f"/api/admin/users/{user_id}/permissions",
                      headers=auth(admin_token),
                      json={"permissions": {SectionKey.DATA_UPLOAD: ALLOW}})

    token = login(upload_client, "dhaka_rm")
    assert upload_client.get("/api/data-upload/types",
                             headers=auth(token)).status_code == 200
    # …and the administration section stays out of reach, because its role
    # ceiling is not something a per-user grant can lift.
    assert upload_client.get("/api/admin/users",
                             headers=auth(token)).status_code == 403


def test_one_user_cannot_commit_another_users_staged_upload(upload_client,
                                                            agent_engine):
    admin_token = login(upload_client)
    created = upload(upload_client, admin_token, "sales",
                     csv_bytes(SALES_HEADERS, sales_rows("UPL-140"))).json()

    with Session(agent_engine) as session:
        user_id = session.execute(
            select(AppUser.user_id).where(AppUser.username == "khulna_rm")
        ).scalar_one()
    upload_client.put(f"/api/admin/users/{user_id}/permissions",
                      headers=auth(admin_token),
                      json={"permissions": {SectionKey.DATA_UPLOAD: ALLOW}})

    other = login(upload_client, "khulna_rm")
    response = commit(upload_client, other, created['upload']['upload_id'])
    assert response.status_code == 403


def test_the_summary_reports_upload_counts(upload_client):
    token = login(upload_client)
    upload(upload_client, token, "sales", csv_bytes(SALES_HEADERS,
                                                    sales_rows("UPL-150")))
    body = upload_client.get("/api/data-upload/summary", headers=auth(token)).json()
    assert body["total_uploads"] >= 1
    assert body["uploads_last_24h"] >= 1
    assert len(body["recent"]) >= 1


def test_a_staged_file_is_removed_once_the_batch_is_terminal(upload_client, tmp_path):
    token = login(upload_client)
    created = upload(upload_client, token, "dim_region", csv_bytes(
        MASTER_HEADERS, [["Z001", "REG960", "Cleanup", "", "", "", ""]])).json()
    assert list(tmp_path.glob("*.csv"))         # staged while awaiting confirmation

    commit(upload_client, token, created['upload']['upload_id'])
    assert not list(tmp_path.glob("*.csv"))     # discarded once imported


def test_upload_errors_are_replaced_not_doubled_on_commit(upload_client,
                                                          agent_engine):
    token = login(upload_client)
    created = upload(upload_client, token, "sales", csv_bytes(SALES_HEADERS, [
        ["2026-08-20", "UPL-160", "SKU001", "TR001", 1, 100, 0, 100],
        ["2026-08-20", "UPL-161", "SKU-NOPE", "TR001", 1, 100, 0, 100],
    ])).json()
    upload_id = created["upload"]["upload_id"]
    commit(upload_client, token, upload_id)

    with Session(agent_engine) as session:
        count = session.execute(
            select(func.count()).select_from(UploadError)
            .where(UploadError.upload_id == upload_id)
        ).scalar_one()
        batch = session.get(UploadBatch, upload_id)
    assert count == 1                       # the validation run's copy was replaced
    assert batch.status == UploadStatus.PARTIAL
    assert batch.stored_path is None


# ---------------------------------------------------------------------------
# Display grouping
# ---------------------------------------------------------------------------


def test_every_upload_type_has_a_display_group():
    """A new upload type must be placed, not silently dropped.

    ``GROUP_BY_KEY`` is hand-written — no property of an upload type says which
    heading a reader expects to find it under — which makes it exactly the kind
    of list this codebase warns about. The Data Management and Data Upload
    screens are the only way data is loaded or corrected, so a type missing from
    the grouping would be a dataset nobody can reach. This is what stops that.
    """
    from app.upload.registry import UPLOAD_TYPES, display_group

    for upload_type in UPLOAD_TYPES:
        assert display_group(upload_type.key), upload_type.key


def test_display_group_refuses_an_unknown_key():
    """It raises rather than defaulting, which is what makes the test above bite."""
    import pytest as _pytest

    from app.upload.registry import display_group

    with _pytest.raises(KeyError):
        display_group("dim_not_a_real_type")


def test_display_groups_and_categories_are_independent():
    """Grouping is presentation; ``category`` stays the processing switch.

    ``upload/service.py`` branches on MASTER vs TRANSACTIONAL and every
    ``upload_batches`` row stores it, so the two must never be conflated: a
    group is free to mix categories and must not change what a category means.
    """
    from app.database.models_admin import UploadCategory
    from app.upload.registry import UPLOAD_TYPES, display_group

    by_key = {t.key: t for t in UPLOAD_TYPES}
    # Market mixes a master dimension with the map's own location table, and
    # Transactions holds exactly the transactional ones.
    assert by_key["sales"].category == UploadCategory.TRANSACTIONAL
    assert display_group("sales") == "TRANSACTIONS"
    assert by_key["dim_company"].category == UploadCategory.MASTER
    assert display_group("dim_company") == "SALES"


def test_both_screens_group_a_record_the_same_way():
    """One record cannot be filed under two headings depending on the page."""
    from app.datamgmt.catalogue import ENTITIES
    from app.upload.registry import display_group

    for entity in ENTITIES:
        assert entity.to_dict()["group"] == display_group(entity.key)
