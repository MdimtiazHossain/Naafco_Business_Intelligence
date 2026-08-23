"""Source readers, header aliasing and the demo generator's safety rule."""

from __future__ import annotations

import subprocess
import sys
from pathlib import Path

import pytest
from openpyxl import Workbook
from sqlalchemy import func, select
from sqlalchemy.orm import Session

from app.database.models_warehouse import FactSales
from app.etl.datasets import SALES, get_dataset, map_headers
from app.etl.pipeline import run_import
from app.etl.readers import CsvSourceReader, ExcelSourceReader, reader_for_file
from conftest_phase2 import sales_row

PROJECT_ROOT = Path(__file__).resolve().parents[2]


# --------------------------------------------------------------------------
# Header mapping
# --------------------------------------------------------------------------


@pytest.mark.parametrize(
    ("header", "expected"),
    [
        ("Invoice No", "invoice_no"),
        ("invoice_number", "invoice_no"),
        ("INVOICE NUMBER", "invoice_no"),
        ("Bill No", "invoice_no"),
        # Both headings reach the same field: what a sales export calls its item
        # column is not what that column resolves against, and a file headed
        # "SKU Code" keeps loading after revision 0022.
        ("SKU Code", "material_code"),
        ("Material Code", "material_code"),
        ("Gross Sales", "gross_sales"),
        ("gross amount", "gross_sales"),
        ("Region Code", "region_code"),
        ("region", "region_code"),
        ("Date", "transaction_date"),
        ("Posting Date", "transaction_date"),
    ],
)
def test_source_headers_map_onto_canonical_fields(header: str, expected: str) -> None:
    assert SALES.resolve_header(header) == expected


def test_unmapped_headers_are_reported(self=None) -> None:
    mapping, unmapped = map_headers(SALES, ["Invoice No", "Mystery Column"])
    assert mapping == {"Invoice No": "invoice_no"}
    assert unmapped == ["Mystery Column"]


def test_business_keys_are_documented_per_dataset() -> None:
    # Two keys, because the source decides which one identifies a line: the
    # ERP's own line number where it supplies one, and the batch-aware
    # combination where it does not. Same invoice, same material, different batch is
    # two lines under either.
    assert get_dataset("sales").business_key_definition == (
        "company_code + invoice_no + invoice_line_no + source_system when "
        "available, otherwise company_code + invoice_no + material_code + "
        "batch_code + source_system"
    )
    # No collection dataset. Receivables left the platform in revision 0020,
    # so there is no reader, no staging table and no business key to document.
    with pytest.raises(ValueError):
        get_dataset("collection")
    # A stock position has no date of its own, so its identity is the place it
    # sits plus the goods it holds: company, plant, storage location and
    # material, and the two dates that belong to the goods rather than to the
    # reporting period.
    #
    # The material code is what makes two materials in the same storage location
    # two positions rather than one overwriting the other. Material group and
    # material brand are *not* in the key: the Material Master records both
    # against the material code, so keying on them would let a mistyped group
    # write a second position for stock that physically exists once.
    assert get_dataset("material_stock").business_key_definition == (
        "company_code + plant_code + storage_location_code + material_code + "
        "production_date + shelf_life_expiration_date + source_system"
    )


def test_unknown_data_type_raises_a_helpful_error() -> None:
    # The message names what *is* supported, which after revision 0020 is the
    # transactional surface entire: sales, material stock and target.
    with pytest.raises(ValueError,
                       match="Supported: sales, material_stock, target"):
        get_dataset("widgets")


# --------------------------------------------------------------------------
# Readers
# --------------------------------------------------------------------------


def test_csv_reader_handles_bom_and_unicode(tmp_path: Path) -> None:
    path = tmp_path / "sales.csv"
    path.write_text(
        "Invoice No,SKU Name\nINV-1,প্রিমিয়াম চা\n", encoding="utf-8-sig"
    )
    reader = CsvSourceReader(path)
    assert reader.headers == ["Invoice No", "SKU Name"]
    rows = list(reader)
    assert rows[0].values["SKU Name"] == "প্রিমিয়াম চা"
    assert rows[0].row_number == 2       # row 1 is the header


def test_csv_reader_sniffs_the_delimiter(tmp_path: Path) -> None:
    path = tmp_path / "sales.txt"
    path.write_text("Invoice No;Quantity\nINV-1;5\n", encoding="utf-8")
    rows = list(CsvSourceReader(path))
    assert rows[0].values["Quantity"] == "5"


def test_csv_reader_skips_blank_lines(tmp_path: Path) -> None:
    path = tmp_path / "sales.csv"
    path.write_text("Invoice No,Quantity\nINV-1,5\n,\nINV-2,7\n", encoding="utf-8")
    assert len(list(CsvSourceReader(path))) == 2


def test_excel_reader_finds_a_header_below_blank_padding(tmp_path: Path) -> None:
    workbook = Workbook()
    worksheet = workbook.active
    worksheet["C4"] = "Invoice No"
    worksheet["D4"] = "Quantity"
    worksheet["C5"] = "INV-1"
    worksheet["D5"] = 5
    path = tmp_path / "sales.xlsx"
    workbook.save(path)

    reader = ExcelSourceReader(path)
    assert reader.headers == ["Invoice No", "Quantity"]
    rows = list(reader)
    assert rows[0].row_number == 5
    assert rows[0].values["Quantity"] == 5


def test_excel_reader_can_target_a_named_sheet(tmp_path: Path) -> None:
    workbook = Workbook()
    workbook.active.title = "Notes"
    sheet = workbook.create_sheet("Sales")
    sheet.append(["Invoice No", "Quantity"])
    sheet.append(["INV-1", 5])
    path = tmp_path / "multi.xlsx"
    workbook.save(path)

    assert list(ExcelSourceReader(path, sheet_name="Sales"))[0].values["Invoice No"] == "INV-1"
    with pytest.raises(ValueError, match="not found"):
        ExcelSourceReader(path, sheet_name="Nope").headers


def test_reader_factory_picks_by_extension(tmp_path: Path) -> None:
    csv_path = tmp_path / "a.csv"
    csv_path.write_text("a,b\n1,2\n", encoding="utf-8")
    assert isinstance(reader_for_file(csv_path), CsvSourceReader)

    workbook = Workbook()
    workbook.active.append(["a"])
    xlsx_path = tmp_path / "a.xlsx"
    workbook.save(xlsx_path)
    assert isinstance(reader_for_file(xlsx_path), ExcelSourceReader)

    with pytest.raises(ValueError, match="Unsupported source file type"):
        reader_for_file(tmp_path / "a.pdf")


def test_missing_file_raises_immediately(tmp_path: Path) -> None:
    with pytest.raises(FileNotFoundError):
        CsvSourceReader(tmp_path / "nope.csv")


# --------------------------------------------------------------------------
# File imports end to end
# --------------------------------------------------------------------------


def test_import_from_a_csv_file(seeded_engine, tmp_path: Path) -> None:
    import csv

    path = tmp_path / "sales.csv"
    row = sales_row()
    with path.open("w", encoding="utf-8", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=list(row))
        writer.writeheader()
        writer.writerow(row)

    result = run_import(seeded_engine, "sales", path, source_system="SAP")
    assert result.status == "COMPLETED", result.error_counts
    with Session(seeded_engine) as session:
        count = session.execute(select(func.count()).select_from(FactSales)).scalar_one()
    assert count == 1


def test_import_from_an_excel_file_with_blank_padding(seeded_engine, tmp_path: Path) -> None:
    row = sales_row()
    workbook = Workbook()
    worksheet = workbook.active
    for column, header in enumerate(row, start=3):
        worksheet.cell(row=3, column=column, value=header)
    for column, value in enumerate(row.values(), start=3):
        worksheet.cell(row=4, column=column, value=value)
    path = tmp_path / "sales.xlsx"
    workbook.save(path)

    result = run_import(seeded_engine, "sales", path)
    assert result.status == "COMPLETED", result.error_counts
    assert result.valid_rows == 1


# --------------------------------------------------------------------------
# Demo generator
# --------------------------------------------------------------------------


def test_demo_generator_refuses_to_invent_master_data(tmp_path: Path, monkeypatch) -> None:
    """With empty master dimensions the generator must stop, not fabricate codes."""
    from alembic import command
    from alembic.config import Config

    url = f"sqlite:///{tmp_path / 'empty.db'}"
    config = Config(str(PROJECT_ROOT / "backend" / "alembic.ini"))
    config.set_main_option(
        "script_location",
        str(PROJECT_ROOT / "backend" / "app" / "database" / "migrations"),
    )
    config.set_main_option("sqlalchemy.url", url)
    monkeypatch.setenv("DATABASE_URL", url)
    command.upgrade(config, "head")

    completed = subprocess.run(
        [sys.executable, str(PROJECT_ROOT / "scripts" / "generate_demo_transactions.py"),
         "--database-url", url, "--out", str(tmp_path / "out")],
        capture_output=True, text=True, encoding="utf-8",
    )

    assert completed.returncode == 2
    assert "hold no records" in completed.stderr
    assert "import_master_data" in completed.stderr
    assert not (tmp_path / "out").exists()


def test_demo_generator_uses_real_master_codes(seeded_engine, tmp_path: Path,
                                               monkeypatch) -> None:
    """With master data present it generates files referencing only real codes."""
    url = str(seeded_engine.url)
    out = tmp_path / "demo"
    completed = subprocess.run(
        [sys.executable, str(PROJECT_ROOT / "scripts" / "generate_demo_transactions.py"),
         "--database-url", url, "--out", str(out), "--format", "csv",
         "--days", "2", "--rows-per-day", "3"],
        capture_output=True, text=True, encoding="utf-8",
    )
    assert completed.returncode == 0, completed.stderr

    sales_file = out / "demo_sales.csv"
    assert sales_file.exists()
    content = sales_file.read_text(encoding="utf-8")
    assert "STR001" in content or "TR001" in content   # real hierarchy codes
    assert "SKU001" in content or "SKU002" in content  # real SKUs
    assert "DEMO" in content

    # And the generated file imports cleanly against the same master data.
    result = run_import(seeded_engine, "sales", sales_file, source_system="DEMO")
    assert result.rejected_rows == 0, result.error_counts
    assert result.valid_rows > 0
