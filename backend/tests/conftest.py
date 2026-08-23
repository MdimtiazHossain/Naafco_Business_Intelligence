"""Shared pytest fixtures.

Tests build their own Excel workbooks in ``tmp_path`` and run the database
against throwaway SQLite files, so the real ``data/Master Data.xlsx`` is only
ever read, never written.
"""

from __future__ import annotations

import sys
from pathlib import Path
from typing import Any, Iterator

import pytest
from openpyxl import Workbook
from sqlalchemy import Engine, create_engine, event

BACKEND_DIR = Path(__file__).resolve().parents[1]
if str(BACKEND_DIR) not in sys.path:
    sys.path.insert(0, str(BACKEND_DIR))

from app.database.models import Base  # noqa: E402

PROJECT_ROOT = BACKEND_DIR.parent


@pytest.fixture(scope="session")
def real_workbook_path() -> Path:
    path = PROJECT_ROOT / "data" / "Master Data.xlsx"
    if not path.exists():
        pytest.skip(f"Master workbook not present at {path}")
    return path


@pytest.fixture
def engine(tmp_path: Path) -> Iterator[Engine]:
    """A SQLite engine with foreign-key enforcement switched on.

    SQLite ignores foreign keys unless ``PRAGMA foreign_keys=ON`` is issued per
    connection; without it the FK tests would pass vacuously.
    """
    engine = create_engine(f"sqlite:///{tmp_path / 'test.db'}", future=True)

    @event.listens_for(engine, "connect")
    def _fk_on(dbapi_connection, _record):  # noqa: ANN001
        cursor = dbapi_connection.cursor()
        cursor.execute("PRAGMA foreign_keys=ON")
        cursor.close()

    Base.metadata.create_all(engine)
    try:
        yield engine
    finally:
        engine.dispose()


def write_sheet(ws, rows: list[list[Any]], start_row: int = 1, start_col: int = 1) -> None:
    """Write a block of values at an arbitrary offset (to simulate blank padding)."""
    for r, row in enumerate(rows):
        for c, value in enumerate(row):
            if value is not None:
                ws.cell(row=start_row + r, column=start_col + c, value=value)


@pytest.fixture
def make_workbook(tmp_path: Path):
    """Factory: ``make_workbook({"Sheet": (rows, start_row, start_col)}) -> Path``."""

    def _make(sheets: dict[str, tuple[list[list[Any]], int, int]],
              name: str = "test.xlsx") -> Path:
        wb = Workbook()
        wb.remove(wb.active)
        for sheet_name, (rows, start_row, start_col) in sheets.items():
            ws = wb.create_sheet(sheet_name)
            write_sheet(ws, rows, start_row, start_col)
        path = tmp_path / name
        wb.save(path)
        return path

    return _make


#: A minimal, internally consistent set of records for the full hierarchy.
HIERARCHY_ROWS: dict[str, list[list[Any]]] = {
    "Company Master": [
        ["Company Code", "Company", "Company Head ID", "Company Head Name"],
        ["C001", "Example Industries Ltd.", "EMP001", "Md. Rahman"],
    ],
    "Business Unit Master": [
        ["Company Code", "BU Code", "BU Name", "BU Head ID", "BU Head Name"],
        ["C001", "BU001", "Consumer Products", "EMP002", "Md. Karim"],
    ],
    "Sales Line Master": [
        ["BU Code", "Sales Line Code", "Sales Line Name", "Sales Line Head ID",
         "Sales Line Head"],
        ["BU001", "SL001", "General Trade", "EMP003", "Md. Alam"],
    ],
    "Zone Master": [
        ["Sales Line Code", "Zone Code", "Zone Name", "Zone Head ID", "Zone Head"],
        ["SL001", "Z001", "Dhaka Zone", "EMP004", "Md. Hasan"],
    ],
    "Region Master": [
        ["Zone Code", "Region Code", "Region", "Region Head ID", "Region Head",
         "Region HQ", "Location"],
        ["Z001", "REG001", "Dhaka", "EMP005", "Md. Selim", "Dhaka", "Dhaka North"],
    ],
    "Area Master": [
        ["Region Code", "Area Code", "Area", "Area Head ID", "Area Head", "Area HQ",
         "Location"],
        ["REG001", "AR001", "Mirpur", "EMP006", "Md. Jamal", "Mirpur", "Mirpur-10"],
    ],
    "Unit Master": [
        ["Area Code", "Unit Code", "Unit", "Unit Head ID", "Unit Head Name", "Unit HQ",
         "Location"],
        ["AR001", "UN001", "Mirpur Unit 1", "EMP007", "Md. Faruk", "Mirpur", "Mirpur-12"],
    ],
    "Territory Master": [
        ["Unit Code", "Territory Code", "Territory", "Territory Head ID", "Territory Head",
         "Territory Head_Phone Number", "Territory HQ", "Location"],
        ["UN001", "TR001", "Kazipara", "EMP008", "Md. Anis", "+8801700000000", "Kazipara",
         "Kazipara Bazar"],
    ],
    "Sub-Territory Master": [
        ["Territory Code", "Sub Territory Code", "Sub Territory", "Sub Territory Head ID",
         "Sub Territory Head", "Sub Territory Head_Phone Number", "Sub Territory HQ",
         "Location"],
        ["TR001", "STR001", "Kazipara North", "EMP009", "Md. Rakib", "+8801800000000",
         "Kazipara", "Kazipara North Block"],
    ],
    "Product Master": [
        ["SKU Code", "SKU Id", "SKU Name En", "SKU Name Bn", "Category", "Product Type",
         "Brand", "Producer Company", "DB Price", "Trade Price", "Retail Price",
         "Pack Size", "Unit Conversation Ratio", "Retailer Unit", "Consumer Unit",
         "Status", "Sales Type", "Total Alt Sku", "Sequence No", "Start Time"],
        ["SKU001", "10001", "Premium Tea 500g", "প্রিমিয়াম চা ৫০০ গ্রাম", "Beverage",
         "Finished Goods", "Example Brand", "Example Industries Ltd.", 120.5, 128, 140,
         "12x500g", 12, "CTN", "PCS", "Active", "Regular", 2, 10, "2024-01-01"],
    ],
}


@pytest.fixture
def full_workbook(make_workbook):
    """A complete, valid workbook: headers on row 2, data from row 3, column B."""

    def _make(overrides: dict[str, list[list[Any]]] | None = None,
              start_row: int = 2, start_col: int = 2) -> Path:
        rows = {name: list(map(list, block)) for name, block in HIERARCHY_ROWS.items()}
        if overrides:
            rows.update({k: list(map(list, v)) for k, v in overrides.items()})
        return make_workbook(
            {name: (block, start_row, start_col) for name, block in rows.items()}
        )

    return _make


# Phase 2 fixtures: a migrated, master-seeded warehouse database.
from conftest_phase2 import (  # noqa: E402,F401
    make_material,
    material_stock_row,
    sales_row,
    seeded_engine,
    target_row,
    warehouse_engine,
)


# Phase 3 fixtures: users with data scopes, loaded facts and an agent.
from conftest_phase3 import (  # noqa: E402,F401
    TODAY,
    agent_engine,
    make_agent,
    make_orchestrator,
    session,
    users,
)
