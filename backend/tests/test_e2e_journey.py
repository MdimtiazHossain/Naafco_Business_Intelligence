"""The critical end-to-end journey from the Phase 4 specification.

Walks the exact sequence a manager performs — sign in, dashboard, filter,
drill down, ask the AI, export, alerts, refusal, sign out — against the real
API. Each step asserts something a user would actually notice.
"""

from __future__ import annotations

import io

import pytest
from fastapi.testclient import TestClient
from sqlalchemy.orm import Session

from app.api.deps import get_session
from app.auth.security import hash_password
from app.database.models_ai import AppUser, AuditLog, Role
from app.main import app
from conftest_phase3 import TODAY

pytest.importorskip("multipart", reason="python-multipart is required by FastAPI forms")

PASSWORD = "Journey-Pass-7"
WINDOW = "date_from=2026-08-01&date_to=2026-08-31"


@pytest.fixture
def journey(agent_engine, users, monkeypatch):
    """A client with passwords set, pinned to the seeded data's calendar."""
    import app.ai.agent as agent_module

    original_init = agent_module.BusinessIntelligenceAgent.__init__

    def pinned_init(self, session, user, llm=None, today=None):
        original_init(self, session, user, llm=llm, today=today or TODAY)

    monkeypatch.setattr(agent_module.BusinessIntelligenceAgent, "__init__", pinned_init)

    with Session(agent_engine) as session:
        for user in session.query(AppUser).all():
            user.password_hash = hash_password(PASSWORD)
        session.commit()

    def _session_override():
        session = Session(bind=agent_engine, expire_on_commit=False, future=True)
        try:
            yield session
        finally:
            session.close()

    app.dependency_overrides[get_session] = _session_override
    try:
        yield TestClient(app)
    finally:
        app.dependency_overrides.clear()


def test_full_management_journey(journey: TestClient, agent_engine) -> None:
    client = journey

    # 1. Sign in as an authorised user.
    login = client.post("/api/auth/login",
                        json={"username": "ceo", "password": PASSWORD, "remember": True})
    assert login.status_code == 200, login.text
    token = login.json()["access_token"]
    headers = {"Authorization": f"Bearer {token}"}
    assert login.json()["user"]["role"] == Role.MANAGEMENT

    profile = client.get("/api/auth/me", headers=headers).json()
    assert profile["scope_description"] == "all regions"

    # 2. Open the dashboard.
    dashboard = client.get(f"/api/dashboard?{WINDOW}", headers=headers).json()
    total_sales = next(k for k in dashboard["kpis"] if k["key"] == "total_sales")
    assert total_sales["value"] == pytest.approx(1_800_000)
    # The frame carries the KPI strip and names its cards; each card is its own
    # request, fetched in parallel by the browser. A journey test walks it the
    # way the reader does rather than reaching for a combined response that no
    # longer exists.
    assert "sales_trend" in dashboard["sections"]
    trend = client.get(f"/api/dashboard/section/sales_trend?{WINDOW}",
                       headers=headers).json()["section"]
    assert trend["rows"]

    # 3. The date range is resolved by the backend, financial year included.
    periods = client.get("/api/period-options", headers=headers).json()
    assert periods["current_financial_year"].startswith("FY ")

    # 4. Choose a region — the filter options are already scoped.
    regions = client.get("/api/master-data/options/region_code", headers=headers).json()
    assert {option["code"] for option in regions["options"]} == {"REG001", "REG002"}

    dhaka_dashboard = client.get(
        f"/api/dashboard?{WINDOW}&region_code=REG001", headers=headers).json()
    dhaka_sales = next(k for k in dhaka_dashboard["kpis"] if k["key"] == "total_sales")
    assert dhaka_sales["value"] == pytest.approx(1_500_000)

    # 5. View sales for that region.
    sales = client.get(f"/api/pages/sales?{WINDOW}&region_code=REG001",
                       headers=headers).json()
    assert sales["summary"]["value"] == pytest.approx(1_500_000)
    assert sales["region_performance"]["rows"]

    # 6. Open performance, and 7. drill down region -> area -> territory.
    performance = client.get(f"/api/pages/performance?level=region&{WINDOW}",
                             headers=headers).json()
    assert performance["next_level"] == "area"

    areas = client.get(
        f"/api/pages/performance?level=area&region_code=REG001&{WINDOW}",
        headers=headers).json()
    assert {row["code"] for row in areas["performance"]["rows"]} == {"AR001"}
    assert areas["next_level"] == "unit"

    territories = client.get(
        f"/api/pages/performance?level=territory&area_code=AR001&{WINDOW}",
        headers=headers).json()
    assert {row["code"] for row in territories["performance"]["rows"]} == {"TR001"}

    # 8-11. Ask the AI, get a grounded answer with data behind it.
    chat = client.post("/api/chat", headers=headers,
                       json={"message": "এই মাসে Dhaka Region-এর sales কত?"}).json()
    assert chat["error_code"] is None, chat["answer"]
    assert chat["intent"] == "SALES_SUMMARY"
    assert chat["filters"]["region_codes"] == ["REG001"]
    assert chat["data"]["value"] == pytest.approx(1_500_000)
    assert chat["conversation_id"]

    # The follow-up inherits the region and metric without repeating them.
    follow_up = client.post(
        "/api/chat", headers=headers,
        json={"message": "Region-wise দেখাও", "conversation_id": chat["conversation_id"]},
    ).json()
    assert follow_up["conversation_id"] == chat["conversation_id"]
    assert follow_up["error_code"] is None

    # A chart is offered for the breakdown, so the UI has something to draw.
    breakdown = client.post("/api/chat", headers=headers,
                            json={"message": "Region-wise sales দেখাও"}).json()
    assert breakdown["chart"] is not None
    assert breakdown["chart"]["data"]

    # 12. Export to Excel.
    from openpyxl import load_workbook

    rows = sales["region_performance"]["rows"]
    excel = client.post("/api/reports/export", headers=headers, json={
        "format": "xlsx", "title": "Sales", "report_name": "Region Sales",
        "date_range": "01 Aug 2026 – 31 Aug 2026",
        "filters": {"region_codes": ["REG001"]}, "rows": rows,
        "kpis": [["Net Sales", "৳15.00 L"]],
    })
    assert excel.status_code == 200
    sheet = load_workbook(io.BytesIO(excel.content)).active
    text = "\n".join(str(c.value) for row in sheet.iter_rows() for c in row if c.value)
    assert "Region Sales" in text and "Generated By" in text

    # 13. Export to PDF.
    pdf = client.post("/api/reports/export", headers=headers, json={
        "format": "pdf", "title": "Sales", "report_name": "Region Sales",
        "rows": rows, "kpis": [["Net Sales", "15.00 L"]],
    })
    assert pdf.status_code == 200
    assert pdf.content.startswith(b"%PDF")

    # 14. Open alerts.
    alerts = client.get(f"/api/alerts?{WINDOW}", headers=headers).json()
    assert alerts["alerts"]
    assert all(alert["link"].startswith("/") for alert in alerts["alerts"])

    # 15. A restricted user cannot reach another region — at any entry point.
    rm_login = client.post("/api/auth/login",
                           json={"username": "dhaka_rm", "password": PASSWORD})
    rm_headers = {"Authorization": f"Bearer {rm_login.json()['access_token']}"}

    refused = client.get(f"/api/pages/sales?{WINDOW}&region_code=REG002",
                         headers=rm_headers)
    assert refused.status_code == 403
    assert "permission" in refused.json()["detail"].lower()

    refused_drill = client.get(
        f"/api/pages/performance?level=area&region_code=REG002&{WINDOW}",
        headers=rm_headers)
    assert refused_drill.status_code == 403

    refused_chat = client.post("/api/chat", headers=rm_headers,
                               json={"message": "Khulna region-এর sales দেখাও"}).json()
    assert refused_chat["error_code"] == "PERMISSION_DENIED"
    assert refused_chat["data"] == {}

    # Their own filter options never even list the other region.
    rm_regions = client.get("/api/master-data/options/region_code",
                            headers=rm_headers).json()
    assert {option["code"] for option in rm_regions["options"]} == {"REG001"}

    # And search cannot be used to discover it.
    assert client.get("/api/master-data/search?q=Khulna",
                      headers=rm_headers).json()["count"] == 0

    # An admin-only page is refused to a non-admin.
    assert client.get("/api/admin/users", headers=headers).status_code == 403

    # 16. Sign out.
    assert client.post("/api/auth/logout", headers=headers).status_code == 200

    # The journey is on the audit trail, and no secret leaked into it.
    with Session(agent_engine) as session:
        actions = {entry.action for entry in session.query(AuditLog).all()}
        detail_text = " ".join(
            str(entry.detail or {}) for entry in session.query(AuditLog).all()
        )
    assert {"LOGIN", "VIEW_REPORT", "EXPORT", "AI_QUERY", "LOGOUT"} <= actions
    assert PASSWORD not in detail_text
    assert "Bearer" not in detail_text


def test_journey_works_on_a_mobile_sized_client(journey: TestClient) -> None:
    """The API is viewport-agnostic; responsiveness is a CSS concern.

    What must hold server-side is that a small client can page through data
    rather than being forced to download everything at once.
    """
    login = journey.post("/api/auth/login",
                         json={"username": "ceo", "password": PASSWORD})
    headers = {"Authorization": f"Bearer {login.json()['access_token']}"}

    page = journey.get(
        f"/api/pages/transactions/sales?{WINDOW}&page=1&page_size=2", headers=headers
    ).json()
    assert len(page["rows"]) <= 2
    assert page["total_pages"] >= 2
