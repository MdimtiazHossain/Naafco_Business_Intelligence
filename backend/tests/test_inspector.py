"""Excel loading, header detection, orientation detection and blank handling."""

from __future__ import annotations

from openpyxl import Workbook

from app.master_data.inspector import (
    detect_header_column,
    detect_header_row,
    inspect_workbook,
)
from app.master_data.schema import EXPECTED_SHEETS


def test_header_is_detected_below_blank_rows_and_columns(full_workbook) -> None:
    """Headers on row 2 / column B must be found without any manual offset."""
    path = full_workbook(start_row=2, start_col=2)
    inspection = inspect_workbook(path)

    company = inspection.sheets["Company Master"].structure
    assert company.orientation == "row_records"
    assert company.header_row == 2
    assert company.blank_rows_before_header == 1
    assert company.blank_columns_before_header == 1
    assert company.columns[:2] == ["Company Code", "Company"]
    assert company.record_count == 1


def test_header_detected_when_there_is_no_blank_padding(full_workbook) -> None:
    path = full_workbook(start_row=1, start_col=1)
    inspection = inspect_workbook(path)
    structure = inspection.sheets["Region Master"].structure
    assert structure.header_row == 1
    assert structure.blank_rows_before_header == 0
    assert structure.record_count == 1


def test_header_detected_deep_into_the_sheet(full_workbook) -> None:
    path = full_workbook(start_row=7, start_col=4)
    inspection = inspect_workbook(path)
    structure = inspection.sheets["Zone Master"].structure
    assert structure.header_row == 7
    assert structure.blank_rows_before_header == 6
    assert structure.blank_columns_before_header == 3
    assert structure.record_count == 1


def test_blank_rows_inside_the_data_block_are_skipped(make_workbook) -> None:
    rows = [
        ["Company Code", "Company", "Company Head ID", "Company Head Name"],
        ["C001", "Alpha Ltd.", "EMP001", "Head One"],
        [None, None, None, None],
        ["C002", "Beta Ltd.", "EMP002", "Head Two"],
    ]
    path = make_workbook({"Company Master": (rows, 2, 2)})
    inspection = inspect_workbook(path)
    sheet = inspection.sheets["Company Master"]

    assert sheet.structure.record_count == 2
    assert sheet.structure.blank_rows_within_data == [4]
    assert [r["Company Code"] for r in sheet.records] == ["C001", "C002"]


def test_vertical_field_list_is_detected_as_column_records(make_workbook) -> None:
    """A transposed sheet with labels down a column and no records."""
    rows = [[label] for label in
            ["Company Code", "Company", "Company Head ID", "Company Head Name"]]
    path = make_workbook({"Company Master": (rows, 2, 3)})
    inspection = inspect_workbook(path)
    structure = inspection.sheets["Company Master"].structure

    assert structure.orientation == "column_records"
    assert structure.header_column == 3
    assert structure.record_count == 0
    assert structure.columns == [
        "Company Code", "Company", "Company Head ID", "Company Head Name"
    ]
    assert any("zero data records" in note for note in structure.notes)


def test_transposed_sheet_with_records_reads_one_record_per_column(make_workbook) -> None:
    rows = [
        ["Company Code", "C001", "C002"],
        ["Company", "Alpha Ltd.", "Beta Ltd."],
        ["Company Head ID", "EMP001", "EMP002"],
        ["Company Head Name", "Head One", "Head Two"],
    ]
    path = make_workbook({"Company Master": (rows, 2, 2)})
    inspection = inspect_workbook(path)
    sheet = inspection.sheets["Company Master"]

    assert sheet.structure.orientation == "column_records"
    assert sheet.structure.record_count == 2
    assert [r["Company Code"] for r in sheet.records] == ["C001", "C002"]
    assert sheet.source_locations == ["column C", "column D"]


def test_ordinary_sheet_is_not_mistaken_for_a_transposed_one(make_workbook) -> None:
    """Many rows must not tip the orientation heuristic towards column layout."""
    header = ["Region Code", "Region", "Region Head ID", "Region Head", "Region HQ"]
    rows = [header] + [
        [f"REG{i:03d}", f"Region {i}", f"EMP{i:04d}", f"Head {i}", f"HQ {i}"]
        for i in range(1, 51)
    ]
    path = make_workbook({"Region Master": (rows, 2, 2)})
    inspection = inspect_workbook(path)
    structure = inspection.sheets["Region Master"].structure

    assert structure.orientation == "row_records"
    assert structure.record_count == 50


def test_styling_only_cells_do_not_create_phantom_records(tmp_path) -> None:
    from openpyxl.styles import PatternFill

    wb = Workbook()
    ws = wb.active
    ws.title = "Company Master"
    ws["B2"] = "Company Code"
    ws["C2"] = "Company"
    ws["B3"] = "C001"
    ws["C3"] = "Alpha Ltd."
    for row in range(4, 60):
        ws.cell(row=row, column=2).fill = PatternFill("solid", fgColor="FFFF00")
    path = tmp_path / "styled.xlsx"
    wb.save(path)

    inspection = inspect_workbook(path)
    assert inspection.sheets["Company Master"].structure.record_count == 1


def test_missing_and_extra_sheets_are_reported_not_deleted(make_workbook) -> None:
    path = make_workbook({
        "Company Master": ([["Company Code", "Company"], ["C001", "Alpha"]], 1, 1),
        "Some Extra Sheet": ([["anything"]], 1, 1),
    })
    inspection = inspect_workbook(path)

    assert "Some Extra Sheet" in inspection.unexpected_sheets
    assert "Some Extra Sheet" in inspection.sheet_names  # kept, never removed
    assert "Territory Master" in inspection.missing_sheets
    # The Product Master is not contracted since revision 0022, so a workbook
    # without it is not incomplete on that account.
    assert "Product Master" not in inspection.missing_sheets
    assert not inspection.is_complete


def test_merged_cells_are_recorded(tmp_path) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Company Master"
    ws["B2"] = "Company Code"
    ws["C2"] = "Company"
    ws["B3"] = "C001"
    ws["C3"] = "Alpha Ltd."
    ws.merge_cells("B1:C1")
    ws["B1"] = "Company Master Data"
    path = tmp_path / "merged.xlsx"
    wb.save(path)

    structure = inspect_workbook(path).sheets["Company Master"].structure
    assert "B1:C1" in structure.merged_ranges
    assert any("merged" in note for note in structure.notes)


def test_duplicate_header_labels_are_suffixed(make_workbook) -> None:
    rows = [["Code", "Location", "Location"], ["C001", "A", "B"]]
    path = make_workbook({"Company Master": (rows, 1, 1)})
    structure = inspect_workbook(path).sheets["Company Master"].structure
    assert structure.columns == ["Code", "Location", "Location__2"]
    assert structure.duplicate_labels == ["Location"]


def test_detect_header_helpers() -> None:
    grid = {(2, 2): "A", (2, 3): "B", (3, 2): 1, (3, 3): 2}
    assert detect_header_row(grid) == 2
    assert detect_header_column(grid) == 2


# --------------------------------------------------------------------------
# The real workbook
# --------------------------------------------------------------------------


def test_real_workbook_has_every_expected_sheet(real_workbook_path) -> None:
    inspection = inspect_workbook(real_workbook_path)
    assert inspection.missing_sheets == []
    assert set(EXPECTED_SHEETS) <= set(inspection.sheet_names)


def test_real_workbook_product_columns(real_workbook_path) -> None:
    inspection = inspect_workbook(real_workbook_path)
    columns = inspection.sheets["Product Master"].structure.columns
    assert columns[0] == "SKU Code"
    assert columns[-1] == "Start Time"
    assert "SKU Name Bn" in columns
    assert len(columns) == 20
