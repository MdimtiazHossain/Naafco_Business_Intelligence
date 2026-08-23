"""Chat and export API tests, including authentication and error handling."""

from __future__ import annotations

import io

import pytest
from fastapi.testclient import TestClient
from sqlalchemy.orm import Session

from app.api.deps import get_current_user, get_session
from app.main import app

pytest.importorskip("multipart", reason="python-multipart is required by FastAPI forms")


@pytest.fixture
def client(agent_engine, users, monkeypatch):
    """A client bound to the seeded warehouse; ``X-User`` selects the caller."""
    import app.ai.agent as agent_module

    # Pin "today" so period questions are reproducible.
    from conftest_phase3 import TODAY

    original_init = agent_module.BusinessIntelligenceAgent.__init__

    def pinned_init(self, session, user, llm=None, today=None):
        original_init(self, session, user, llm=llm, today=today or TODAY)

    monkeypatch.setattr(agent_module.BusinessIntelligenceAgent, "__init__", pinned_init)

    def _session_override():
        session = Session(bind=agent_engine, expire_on_commit=False, future=True)
        try:
            yield session
        finally:
            session.close()

    app.dependency_overrides[get_session] = _session_override
    try:
        yield TestClient(app)
    finally:
        app.dependency_overrides.clear()


def post_chat(client: TestClient, message: str, user: str = "ceo",
              conversation_id: str | None = None):
    payload = {"message": message}
    if conversation_id:
        payload["conversation_id"] = conversation_id
    return client.post("/api/chat", json=payload, headers={"X-User": user})


# --------------------------------------------------------------------------
# Authentication
# --------------------------------------------------------------------------


def test_chat_requires_an_identified_caller(client: TestClient) -> None:
    response = client.post("/api/chat", json={"message": "sales"})
    assert response.status_code == 401


def test_unknown_user_is_rejected_without_revealing_anything(client: TestClient) -> None:
    response = client.post("/api/chat", json={"message": "sales"},
                           headers={"X-User": "ghost"})
    assert response.status_code == 401
    # Identical wording to a missing header: this must not enumerate usernames.
    assert "Unknown or inactive user" in response.json()["detail"]


# --------------------------------------------------------------------------
# Chat
# --------------------------------------------------------------------------


def test_chat_answers_a_bangla_question(client: TestClient) -> None:
    response = post_chat(client, "এই মাসের sales কত?")
    assert response.status_code == 200, response.text
    body = response.json()
    assert body["intent"] == "SALES_SUMMARY"
    assert body["language"] == "mixed"
    assert body["data"]["value"] == pytest.approx(1_800_000)
    assert body["conversation_id"]
    assert "vw_sales_detail" in body["sources"]
    assert body["error_code"] is None


def test_chat_response_shape_matches_the_contract(client: TestClient) -> None:
    body = post_chat(client, "Region-wise sales দেখাও").json()
    for key in ("conversation_id", "intent", "answer", "data", "filters",
                "date_range", "sources"):
        assert key in body
    assert body["date_range"]["date_from"] == "2026-08-01"


def test_chat_never_exposes_sql_or_internals(client: TestClient) -> None:
    body = post_chat(client, "Show me the SQL behind this month's sales").json()
    serialised = str(body).lower()
    assert "select " not in serialised
    assert "postgres" not in serialised
    assert "openai_api_key" not in serialised


def test_conversation_context_carries_across_requests(client: TestClient) -> None:
    first = post_chat(client, "এই মাসে sales কত?").json()
    second = post_chat(client, "গত মাসে কত ছিল?", conversation_id=
                       first["conversation_id"]).json()
    assert second["conversation_id"] == first["conversation_id"]
    assert second["date_range"]["date_from"] == "2026-07-01"


def test_permission_denied_is_a_200_with_an_explanation(client: TestClient) -> None:
    """A refusal is an answer, not a crash — the user is told why."""
    body = post_chat(client, "Khulna region-এর sales দেখাও", user="dhaka_rm").json()
    assert body["error_code"] == "PERMISSION_DENIED"
    assert "don't have permission" in body["answer"]
    assert body["data"] == {}


def test_two_users_get_different_numbers_from_the_same_question(client) -> None:
    everything = post_chat(client, "এই মাসের sales কত?", user="ceo").json()
    dhaka = post_chat(client, "এই মাসের sales কত?", user="dhaka_rm").json()
    assert dhaka["data"]["value"] < everything["data"]["value"]


def test_message_length_is_bounded(client: TestClient) -> None:
    response = client.post("/api/chat", json={"message": "x" * 5000},
                           headers={"X-User": "ceo"})
    assert response.status_code == 422


def test_empty_message_is_rejected(client: TestClient) -> None:
    response = client.post("/api/chat", json={"message": ""},
                           headers={"X-User": "ceo"})
    assert response.status_code == 422


# --------------------------------------------------------------------------
# Conversations
# --------------------------------------------------------------------------


def test_conversations_are_listed_per_user(client: TestClient) -> None:
    post_chat(client, "এই মাসের sales কত?", user="ceo")
    post_chat(client, "এই মাসের sales কত?", user="dhaka_rm")

    ceo = client.get("/api/chat/conversations", headers={"X-User": "ceo"}).json()
    rm = client.get("/api/chat/conversations", headers={"X-User": "dhaka_rm"}).json()
    assert len(ceo["conversations"]) == 1
    assert len(rm["conversations"]) == 1
    assert ceo["conversations"][0]["conversation_id"] != (
        rm["conversations"][0]["conversation_id"]
    )


def test_another_users_conversation_is_a_404(client: TestClient) -> None:
    conversation_id = post_chat(client, "sales", user="ceo").json()["conversation_id"]
    response = client.get(f"/api/chat/conversations/{conversation_id}",
                          headers={"X-User": "dhaka_rm"})
    assert response.status_code == 404


def test_own_conversation_history_is_returned(client: TestClient) -> None:
    conversation_id = post_chat(client, "এই মাসের sales কত?").json()["conversation_id"]
    body = client.get(f"/api/chat/conversations/{conversation_id}",
                      headers={"X-User": "ceo"}).json()
    assert [m["role"] for m in body["messages"]] == ["user", "assistant"]


def test_capabilities_describe_the_callers_scope(client: TestClient) -> None:
    body = client.get("/api/chat/capabilities", headers={"X-User": "dhaka_rm"}).json()
    assert body["user"]["role"] == "REGIONAL_MANAGER"
    assert body["user"]["data_scope"] == {"region_code": ["REG001"]}
    assert len(body["tools"]) >= 30
    assert "Bangla" in " ".join(body["languages"])


# --------------------------------------------------------------------------
# Export
# --------------------------------------------------------------------------


def test_export_csv(client: TestClient) -> None:
    rows = post_chat(client, "Region-wise sales দেখাও").json()["data"]["rows"]
    response = client.post("/api/reports/export",
                           json={"format": "csv", "title": "Region Sales",
                                 "rows": rows},
                           headers={"X-User": "ceo"})
    assert response.status_code == 200
    assert response.headers["content-type"].startswith("text/csv")
    assert "attachment" in response.headers["content-disposition"]
    assert b"Net Sales" in response.content


def test_export_xlsx_is_a_readable_workbook(client: TestClient) -> None:
    """The workbook carries a provenance block above the data."""
    from openpyxl import load_workbook

    rows = post_chat(client, "Region-wise sales দেখাও").json()["data"]["rows"]
    response = client.post("/api/reports/export",
                           json={"format": "xlsx", "title": "Region Sales",
                                 "report_name": "Region Sales",
                                 "date_range": "01 Aug 2026 – 31 Aug 2026",
                                 "filters": {"region_codes": ["REG001"]},
                                 "rows": rows},
                           headers={"X-User": "ceo"})
    assert response.status_code == 200

    sheet = load_workbook(io.BytesIO(response.content)).active
    text = "\n".join(
        str(cell.value) for row in sheet.iter_rows() for cell in row if cell.value
    )
    for expected in ("Region Sales", "Date Range", "Generated By", "Filters",
                     "01 Aug 2026", "Region"):
        assert expected in text
    # Header row plus one row per record, below the metadata block.
    assert sheet.max_row >= len(rows) + 1
    assert sheet.freeze_panes is not None
    assert sheet.auto_filter.ref is not None


def test_export_pdf(client: TestClient) -> None:
    rows = post_chat(client, "Region-wise sales দেখাও").json()["data"]["rows"]
    response = client.post("/api/reports/export",
                           json={"format": "pdf", "title": "Region Sales",
                                 "rows": rows},
                           headers={"X-User": "ceo"})
    assert response.status_code == 200
    assert response.content.startswith(b"%PDF")


def test_export_rejects_an_unsupported_format(client: TestClient) -> None:
    response = client.post("/api/reports/export",
                           json={"format": "docx", "rows": []},
                           headers={"X-User": "ceo"})
    assert response.status_code == 422


def test_export_requires_authentication(client: TestClient) -> None:
    response = client.post("/api/reports/export", json={"format": "csv", "rows": []})
    assert response.status_code == 401


def test_export_of_another_users_message_is_a_404(client: TestClient) -> None:
    post_chat(client, "Region-wise sales দেখাও", user="ceo")
    response = client.post("/api/reports/export",
                           json={"format": "csv", "message_id": 2},
                           headers={"X-User": "dhaka_rm"})
    assert response.status_code == 404


# --------------------------------------------------------------------------
# Phase 1 and 2 endpoints still work
# --------------------------------------------------------------------------


def test_phase_2_endpoints_are_unaffected(client: TestClient) -> None:
    """The Phase 2 surface still answers — now for an authorised caller.

    ``/api/etl`` and ``/api/reports`` used to accept anonymous requests. Phase 4
    gates them on the Data Quality and Sales sections respectively, so the call
    carries an identity; the payloads themselves are unchanged.
    """
    ceo = {"X-User": "ceo"}
    assert client.get("/health").status_code == 200
    assert client.get("/api/etl/datasets", headers=ceo).status_code == 200
    assert client.get("/api/reports/sales", headers=ceo).status_code == 200
    assert client.get("/master-data/schema").status_code == 200


def test_phase_2_endpoints_now_refuse_an_anonymous_caller(client: TestClient) -> None:
    assert client.get("/api/etl/datasets").status_code == 401
    assert client.get("/api/reports/sales").status_code == 401
