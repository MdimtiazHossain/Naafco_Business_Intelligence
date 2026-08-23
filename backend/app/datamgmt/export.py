"""Exporting a table.

An export is a read, so it goes through exactly the query the table went
through: same scope, same filters, same search. What it must never be is a way
round the pagination — asking for a CSV of a table you may only see one region
of has to produce that region, and nothing else.

Two consequences of that, both deliberate:

* the export re-runs the list query rather than accepting rows from the client,
  so there is no path by which the browser can widen what it receives;
* it is bounded. A file of every fact row is not an export, it is a database
  copy; past the cap the caller is told to narrow the filters, because a
  truncated file that does not say so is worse than a refusal.
"""

from __future__ import annotations

import csv
import io
from typing import Any

from sqlalchemy.orm import Session

from ..ai.permission_filter import UserContext
from .catalogue import ManagedEntity
from .query import ListRequest, list_master, list_transactions

#: The most rows one export may contain.
MAX_EXPORT_ROWS = 20_000

#: Rows fetched per query while streaming a large export together.
CHUNK = 2_000

CSV_MEDIA_TYPE = "text/csv; charset=utf-8"
XLSX_MEDIA_TYPE = (
    "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
)


class ExportTooLarge(Exception):
    """More rows match than an export may carry."""


def collect(session: Session, user: UserContext, entity: ManagedEntity,
            request: ListRequest, columns: list[str] | None = None,
            ) -> tuple[list[str], list[dict[str, Any]]]:
    """Every row the caller's filters select, in pages, up to the cap."""
    lister = list_master if entity.is_master else list_transactions

    probe = lister(session, user, entity,
                   _page(request, page=1, size=1))
    if probe.total > MAX_EXPORT_ROWS:
        raise ExportTooLarge(
            f"{probe.total:,} records match. An export is limited to "
            f"{MAX_EXPORT_ROWS:,} rows — narrow the filters or the date range, "
            "or export one page at a time."
        )

    chosen = _columns(entity, columns, probe.columns)
    rows: list[dict[str, Any]] = []
    page = 1
    while len(rows) < probe.total:
        result = lister(session, user, entity, _page(request, page=page,
                                                     size=CHUNK))
        if not result.rows:
            break
        rows.extend(result.rows)
        page += 1
    return chosen, rows


def _page(request: ListRequest, *, page: int, size: int) -> ListRequest:
    from dataclasses import replace

    return replace(request, page=page, page_size=size)


def _columns(entity: ManagedEntity, requested: list[str] | None,
             available: list[str]) -> list[str]:
    """Which columns to write.

    The caller may pass the columns their table is currently showing, so the
    file matches the screen. Anything not on the entity is dropped rather than
    trusted — a column name from the client must never reach a row lookup it
    could not otherwise reach.
    """
    known = {f.name for f in entity.fields} | set(available)
    if requested:
        chosen = [name for name in requested if name in known]
        if chosen:
            return chosen
    return [name for name in available if not name.startswith("_")]


def _cell(value: Any) -> Any:
    if value is None:
        return ""
    if isinstance(value, bool):
        return "Yes" if value else "No"
    return value


def to_csv(entity: ManagedEntity, columns: list[str],
           rows: list[dict[str, Any]]) -> bytes:
    """A CSV with a header row of human labels.

    UTF-8 with a BOM, because Excel on Windows reads a BOM-less UTF-8 CSV as
    the system code page and mangles every Bangla name in it.
    """
    labels = entity.field_by_name
    buffer = io.StringIO(newline="")
    writer = csv.writer(buffer, lineterminator="\r\n")
    writer.writerow([
        labels[name].label if name in labels else name.replace("_", " ").title()
        for name in columns
    ])
    for row in rows:
        writer.writerow([_cell(row.get(name)) for name in columns])
    return b"\xef\xbb\xbf" + buffer.getvalue().encode("utf-8")


def to_xlsx(entity: ManagedEntity, columns: list[str],
            rows: list[dict[str, Any]]) -> bytes:
    """The same content as a worksheet, when openpyxl is installed."""
    from openpyxl import Workbook
    from openpyxl.styles import Font
    from openpyxl.utils import get_column_letter

    labels = entity.field_by_name
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = entity.label[:31] or "Data"

    headers = [
        labels[name].label if name in labels else name.replace("_", " ").title()
        for name in columns
    ]
    sheet.append(headers)
    for cell in sheet[1]:
        cell.font = Font(bold=True)
    sheet.freeze_panes = "A2"

    for row in rows:
        sheet.append([_cell(row.get(name)) for name in columns])

    for index, header in enumerate(headers, start=1):
        width = max(len(str(header)) + 2, 12)
        sheet.column_dimensions[get_column_letter(index)].width = min(width, 40)

    stream = io.BytesIO()
    workbook.save(stream)
    return stream.getvalue()


def filename(entity: ManagedEntity, extension: str) -> str:
    import datetime as dt

    stamp = dt.date.today().isoformat()
    slug = entity.key.replace("dim_", "").replace("_", "-")
    return f"{slug}-{stamp}.{extension}"


__all__ = [
    "MAX_EXPORT_ROWS",
    "CSV_MEDIA_TYPE",
    "XLSX_MEDIA_TYPE",
    "ExportTooLarge",
    "collect",
    "to_csv",
    "to_xlsx",
    "filename",
]
