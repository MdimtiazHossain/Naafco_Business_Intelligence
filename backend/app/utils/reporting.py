"""Report writers: profiling JSON/Excel, data dictionary and ER diagram."""

from __future__ import annotations

import json
from datetime import datetime, timezone
from decimal import Decimal
from pathlib import Path
from typing import Any

from ..master_data.profiler import WorkbookProfile
from ..master_data.schema import (
    HIERARCHY,
    SQL_TYPE_BY_KIND,
    TABLE_SPECS,
    VALIDATION_RULES,
    TableSpec,
)


def _json_default(obj: Any) -> Any:
    if isinstance(obj, Decimal):
        return str(obj)
    if isinstance(obj, datetime):
        return obj.isoformat()
    if isinstance(obj, Path):
        return str(obj)
    raise TypeError(f"Not JSON serialisable: {type(obj)!r}")


def write_json(path: Path, payload: dict[str, Any]) -> Path:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(
        json.dumps(payload, indent=2, ensure_ascii=False, default=_json_default),
        encoding="utf-8",
    )
    return path


# ---------------------------------------------------------------------------
# Profiling workbook
# ---------------------------------------------------------------------------


def write_profile_workbook(path: Path, profile: WorkbookProfile) -> Path:
    """Write the profiling report as an Excel workbook (overview + per sheet)."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    path.parent.mkdir(parents=True, exist_ok=True)
    wb = Workbook()
    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill("solid", fgColor="1F4E79")

    def style_header(ws, ncols: int) -> None:
        for c in range(1, ncols + 1):
            cell = ws.cell(row=1, column=c)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = Alignment(vertical="center", wrap_text=True)
        ws.freeze_panes = "A2"

    def autosize(ws, ncols: int, max_width: int = 46) -> None:
        for c in range(1, ncols + 1):
            longest = max(
                (len(str(ws.cell(row=r, column=c).value or "")) for r in range(1, ws.max_row + 1)),
                default=10,
            )
            ws.column_dimensions[get_column_letter(c)].width = min(max(12, longest + 2), max_width)

    overview = wb.active
    overview.title = "Overview"
    overview_cols = [
        "Sheet", "Target Table", "Orientation", "Header Row", "Header Column",
        "Records", "Columns", "Fully Duplicated Rows", "Missing Expected Fields",
        "Unexpected Fields", "Notes",
    ]
    overview.append(overview_cols)
    for sheet in profile.sheets:
        overview.append([
            sheet.sheet_name,
            sheet.table_name or "-",
            sheet.orientation,
            sheet.header_row if sheet.header_row is not None else "-",
            sheet.header_column if sheet.header_column is not None else "-",
            sheet.record_count,
            sheet.column_count,
            sheet.fully_duplicated_rows,
            ", ".join(sheet.missing_expected_fields) or "-",
            ", ".join(sheet.unexpected_fields) or "-",
            " | ".join(sheet.notes) or "-",
        ])
    style_header(overview, len(overview_cols))
    autosize(overview, len(overview_cols))

    column_cols = [
        "Column", "Warehouse Column", "Inferred Type", "Declared SQL Type",
        "Records", "Non-null", "Null", "Null %", "Unique", "Duplicates",
        "Min Len", "Max Len", "Whitespace Padded", "Empty Strings",
        "Numeric-typed Cells", "Example Values",
    ]

    all_columns = wb.create_sheet("All Columns")
    all_columns.append(["Sheet", *column_cols])

    for sheet in profile.sheets:
        title = sheet.sheet_name[:31]
        ws = wb.create_sheet(title)
        ws.append(column_cols)
        for col in sheet.columns:
            row = [
                col.column_name, col.warehouse_column or "-", col.inferred_type,
                col.declared_type or "-", col.total_records, col.non_null_records,
                col.null_records, col.null_percent, col.unique_count, col.duplicate_count,
                col.min_length if col.min_length is not None else "-",
                col.max_length if col.max_length is not None else "-",
                col.whitespace_padded, col.empty_string, col.numeric_typed,
                ", ".join(col.example_values) or "-",
            ]
            ws.append(row)
            all_columns.append([sheet.sheet_name, *row])
        style_header(ws, len(column_cols))
        autosize(ws, len(column_cols))

    style_header(all_columns, len(column_cols) + 1)
    autosize(all_columns, len(column_cols) + 1)

    wb.save(path)
    return path


# ---------------------------------------------------------------------------
# Mermaid ER diagram
# ---------------------------------------------------------------------------


def _mermaid_type(sql_type: str) -> str:
    """Mermaid ER attribute types may not contain brackets or spaces."""
    return sql_type.split("(")[0].strip().lower().replace(" ", "_")


def build_er_diagram(specs: tuple[TableSpec, ...] = TABLE_SPECS) -> str:
    """Render the dimension model as a Mermaid ``erDiagram``."""
    lines = ["erDiagram"]
    for spec in specs:
        if spec.has_foreign_key():
            parent = next(s for s in specs if s.table == spec.parent_table)
            lines.append(
                f"    {parent.table.upper()} ||--o{{ {spec.table.upper()} : contains"
            )
    lines.append("")
    for spec in specs:
        lines.append(f"    {spec.table.upper()} {{")
        lines.append(f"        bigint {spec.surrogate_key} PK")
        for col in spec.columns:
            marker = ""
            if col.column == spec.business_key:
                marker = " UK"
            elif spec.has_foreign_key() and col.column == spec.parent_column:
                marker = " FK"
            null_marker = "" if col.nullable else ' "not null"'
            lines.append(
                f"        {_mermaid_type(col.sql_type)} {col.column}{marker}{null_marker}"
            )
        lines.append("        timestamp created_at")
        lines.append("        timestamp updated_at")
        lines.append("    }")
        lines.append("")
    return "\n".join(lines).rstrip() + "\n"


# ---------------------------------------------------------------------------
# Data dictionary
# ---------------------------------------------------------------------------


def build_data_dictionary(profile: WorkbookProfile,
                          specs: tuple[TableSpec, ...] = TABLE_SPECS) -> str:
    """Render the Markdown data dictionary for every dimension table."""
    profile_by_sheet = {p.sheet_name: p for p in profile.sheets}
    generated = datetime.now(timezone.utc).isoformat(timespec="seconds")

    out: list[str] = []
    out.append("# Master Data Dictionary")
    out.append("")
    out.append(f"- **Source workbook:** `{profile.source_path}`")
    out.append(f"- **Generated (UTC):** {generated}")
    out.append(f"- **Sheets inspected:** {len(profile.sheets)}")
    out.append(f"- **Total source records:** {profile.total_records}")
    out.append("")
    out.append("This document is generated by `scripts/inspect_master_data.py`. Every source "
               "field listed here was read from the workbook; nothing is assumed.")
    out.append("")

    out.append("## Organisational hierarchy")
    out.append("")
    out.append("```")
    indent = 0
    for table in HIERARCHY:
        spec = next(s for s in specs if s.table == table)
        out.append(" " * indent + spec.sheet.replace(" Master", ""))
        indent += 4
        if table != HIERARCHY[-1]:
            out.append(" " * indent + "|")
    out.append("```")
    out.append("")

    out.append("## Entity relationship diagram")
    out.append("")
    out.append("```mermaid")
    out.append(build_er_diagram(specs).rstrip())
    out.append("```")
    out.append("")

    out.append("## Table summary")
    out.append("")
    out.append("| Table | Source Sheet | Surrogate Key | Business Key | Parent | Source Records |")
    out.append("|---|---|---|---|---|---|")
    for spec in specs:
        sheet_profile = profile_by_sheet.get(spec.sheet)
        parent = (
            f"`{spec.parent_table}.{spec.parent_key}`" if spec.has_foreign_key() else "—"
        )
        out.append(
            f"| `{spec.table}` | {spec.sheet} | `{spec.surrogate_key}` | "
            f"`{spec.business_key}` | {parent} | "
            f"{sheet_profile.record_count if sheet_profile else 0} |"
        )
    out.append("")

    for spec in specs:
        sheet_profile = profile_by_sheet.get(spec.sheet)
        out.append(f"## `{spec.table}`")
        out.append("")
        out.append(f"- **Purpose:** {spec.purpose}")
        out.append(f"- **Source sheet:** `{spec.sheet}`")
        out.append(f"- **Primary key (surrogate):** `{spec.surrogate_key}` (BIGSERIAL)")
        out.append(f"- **Business key:** `{spec.business_key}` (UNIQUE, official code, TEXT)")
        if spec.has_foreign_key():
            out.append(
                f"- **Foreign key:** `{spec.table}.{spec.parent_column}` → "
                f"`{spec.parent_table}.{spec.parent_key}`"
            )
        else:
            out.append("- **Foreign keys:** none (independent dimension)")
        out.append(f"- **Indexes:** " + ", ".join(f"`{c}`" for c in spec.indexed_columns))
        if sheet_profile:
            out.append(
                f"- **Source layout:** {sheet_profile.orientation}, "
                + (f"header row {sheet_profile.header_row}"
                   if sheet_profile.header_row else
                   f"header column {sheet_profile.header_column}")
                + f", {sheet_profile.record_count} record(s)"
            )
        out.append("")
        out.append("| Column | SQL Type | Nullable | Source Sheet | Source Field | "
                   "Relationship | Description | Example |")
        out.append("|---|---|---|---|---|---|---|---|")
        out.append(
            f"| `{spec.surrogate_key}` | BIGSERIAL | NO | — | — | PK | "
            "Internal surrogate key. Never exposed as a business identifier. | 1 |"
        )
        for col in spec.columns:
            if col.column == spec.business_key:
                rel = "Business key (UNIQUE)"
            elif spec.has_foreign_key() and col.column == spec.parent_column:
                rel = f"FK → `{spec.parent_table}.{spec.parent_key}`"
            else:
                rel = "—"
            out.append(
                f"| `{col.column}` | {col.sql_type} | {'YES' if col.nullable else 'NO'} | "
                f"{spec.sheet} | {col.source_field} | {rel} | {col.description} | "
                f"{col.example} |"
            )
        out.append(
            "| `created_at` | TIMESTAMPTZ | NO | — | — | — | Row insert timestamp. | now() |"
        )
        out.append(
            "| `updated_at` | TIMESTAMPTZ | NO | — | — | — | Row update timestamp. | now() |"
        )
        out.append("")

        if sheet_profile and sheet_profile.columns:
            out.append("### Source column profile")
            out.append("")
            out.append("| Source Column | Inferred Type | Records | Non-null | Null | Null % | "
                       "Unique | Duplicates | Examples |")
            out.append("|---|---|---|---|---|---|---|---|---|")
            for col in sheet_profile.columns:
                examples = ", ".join(col.example_values) or "—"
                out.append(
                    f"| {col.column_name} | {col.inferred_type} | {col.total_records} | "
                    f"{col.non_null_records} | {col.null_records} | {col.null_percent} | "
                    f"{col.unique_count} | {col.duplicate_count} | {examples} |"
                )
            out.append("")
        if sheet_profile and sheet_profile.notes:
            out.append("### Notes")
            out.append("")
            for note in sheet_profile.notes:
                out.append(f"- {note}")
            out.append("")

    out.append("## Validation rules")
    out.append("")
    out.append("| Rule | Name | Severity | Applies To | Description |")
    out.append("|---|---|---|---|---|")
    for rule in VALIDATION_RULES:
        out.append(
            f"| {rule.rule_id} | {rule.name} | {rule.severity} | {rule.applies_to} | "
            f"{rule.description} |"
        )
    out.append("")

    out.append("## Type mapping")
    out.append("")
    out.append("| Field kind | PostgreSQL type | Handling |")
    out.append("|---|---|---|")
    handling = {
        "code": "Official business code. Always TEXT — `001` never becomes `1`.",
        "text": "Free text; Unicode (Bangla) preserved, only outer whitespace trimmed.",
        "phone": "Phone number kept as TEXT so leading zeros and `+` survive.",
        "integer": "Whole number; coercion failure is a blocking validation error.",
        "decimal": "Money/ratio at 4 decimal places.",
        "timestamp": "Parsed from Excel datetimes or common string formats.",
    }
    for kind, sql in SQL_TYPE_BY_KIND.items():
        out.append(f"| {kind.value} | {sql} | {handling.get(kind.value, '')} |")
    out.append("")

    return "\n".join(out) + "\n"


def write_text(path: Path, content: str) -> Path:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(content, encoding="utf-8")
    return path


__all__ = [
    "write_json",
    "write_text",
    "write_profile_workbook",
    "build_er_diagram",
    "build_data_dictionary",
]
